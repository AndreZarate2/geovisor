#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build DOCX/PDF technical reports for drawn concession/convenio zones."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import unicodedata
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(20, 32, 52)
MUTED = RGBColor(84, 99, 122)
FILL = "E8EEF5"
ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "img" / "cvc-logo.png"
MAP_DIR = ROOT / "reports" / "_maps"
BODY_FONT = "Arial"
BODY_SIZE = 11
LINE_SPACING = 1.15


def val(obj: dict[str, Any], path: str, default: Any = "") -> Any:
    cur: Any = obj
    for part in path.split("."):
        if not isinstance(cur, dict) or part not in cur:
            return default
        cur = cur[part]
    return cur


def fmt(n: Any, digits: int = 2, dash: str = "-") -> str:
    try:
        if n is None or n == "":
            return dash
        return f"{float(n):,.{digits}f}"
    except Exception:
        return str(n)


def clean_ai_text(value: Any) -> str:
    text_value = str(value or "")
    text_value = re.sub(r"\*\*([^*]+)\*\*", r"\1", text_value)
    text_value = re.sub(r"__([^_]+)__", r"\1", text_value)
    text_value = re.sub(r"`{1,3}", "", text_value)
    text_value = re.sub(r"^\s{0,3}#{1,6}\s*", "", text_value, flags=re.MULTILINE)
    text_value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text_value)
    return text_value.strip()


def clean_pdf_text(value: Any, limit: int = 420) -> str:
    text_value = clean_ai_text(value)
    text_value = re.sub(r"\s+", " ", text_value).strip()
    if len(text_value) > limit:
        return text_value[: max(0, limit - 1)].rstrip() + "..."
    return text_value


def text(v: Any, fallback: str = "-") -> str:
    if v is None or v == "":
        return fallback
    return str(v)


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except Exception:
        return default


def template_active(meta: dict[str, Any]) -> bool:
    return Path(str(meta.get("template") or "")).is_file()


def template_preserve_body(meta: dict[str, Any]) -> bool:
    return bool(meta.get("preserve_template_body") or meta.get("keep_template_body"))


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    paragraph.text = new_text
    format_paragraph(paragraph)


def _template_paragraph_text(current: str, meta: dict[str, Any]) -> str | None:
    raw = str(current or "")
    compact = re.sub(r"\s+", " ", raw.strip())
    upper = compact.upper()
    if not compact:
        return None
    project = str(meta.get("project") or "Evaluacion tecnica")
    author = str(meta.get("author") or "CVC Energia")
    generated = str(meta.get("generated_at") or datetime.now().strftime("%Y-%m-%d"))
    if upper.startswith("RESUMEN EJECUTIVO: PROYECTO"):
        return f"Resumen Ejecutivo: Proyecto \"{project}\""
    if upper.startswith("PROYECTO:"):
        return f"PROYECTO: \"{project}\""
    if upper.startswith("PREPARADO POR:"):
        return f"Preparado por: {author}"
    if "MES ACTUAL" in upper:
        return generated
    if "INDICE MODIFICABLE" in upper:
        return "INDICE"
    if "DE AQUI PARA ABAJO" in upper or "DE AQUÍ PARA ABAJO" in upper:
        return "DESARROLLO DEL INFORME"
    return None


def apply_template_fields(doc: Document, meta: dict[str, Any]) -> None:
    containers = list(doc.paragraphs)
    for section in doc.sections:
        containers.extend(section.header.paragraphs)
        containers.extend(section.footer.paragraphs)
    for paragraph in containers:
        new_text = _template_paragraph_text(paragraph.text, meta)
        if new_text is not None:
            _replace_paragraph_text(paragraph, new_text)


def format_run(run, size: float = BODY_SIZE, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = BODY_FONT
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = color or DARK


def format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float = BODY_SIZE) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        format_run(run, size=size)


def add_body_paragraph(doc: Document, content: Any) -> None:
    p = doc.add_paragraph(str(content or ""))
    format_paragraph(p)


def geometry_points(data: dict[str, Any]) -> list[list[float]]:
    raw = data.get("polygon")
    if not isinstance(raw, list):
        raw = val(data, "geometry.coordinates", [])
    pts: list[list[float]] = []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                lat, lng = item.get("lat"), item.get("lng")
            elif isinstance(item, (list, tuple)) and len(item) >= 2:
                lat, lng = item[0], item[1]
            else:
                continue
            try:
                pts.append([float(lat), float(lng)])
            except Exception:
                continue
    return pts


REFERENCE_LAYERS = [
    ("MIDAGRI_Departamentos.geojson", "polygon", "#94A3B8", None, 1, 60),
    ("Concesiones_Distribucion_vigente_definitiva.geojson", "polygon", "#16A34A", "#DCFCE7", 2, 80),
    ("Concesiones_Distribucion_vigente_rural.geojson", "polygon", "#0891B2", "#CFFAFE", 2, 80),
    ("Concesiones_Distribucion_tramite_definitiva.geojson", "polygon", "#F59E0B", "#FEF3C7", 2, 80),
    ("Concesiones_Distribucion_tramite_rural.geojson", "polygon", "#A855F7", "#F3E8FF", 2, 80),
    ("LT_Operacionm.geojson", "line", "#F97316", None, 2, 160),
    ("LT_Proyectadas.geojson", "line", "#7C3AED", None, 2, 120),
]


def _geojson_path(name: str) -> Path:
    return ROOT / "data" / "geojson" / name


def _coord_lines(geometry: dict[str, Any], limit: int = 600) -> list[list[list[float]]]:
    gtype = str(geometry.get("type") or "")
    coords = geometry.get("coordinates")
    lines: list[list[list[float]]] = []

    def point(pair: Any) -> list[float] | None:
        if isinstance(pair, (list, tuple)) and len(pair) >= 2:
            try:
                return [float(pair[1]), float(pair[0])]
            except Exception:
                return None
        return None

    def simple_line(raw: Any) -> list[list[float]]:
        out: list[list[float]] = []
        if not isinstance(raw, list):
            return out
        step = max(1, int(len(raw) / max(1, limit)))
        for item in raw[::step]:
            p = point(item)
            if p:
                out.append(p)
        return out

    if gtype == "Point":
        p = point(coords)
        if p:
            lines.append([p])
    elif gtype in {"LineString", "MultiPoint"}:
        line = simple_line(coords)
        if line:
            lines.append(line)
    elif gtype in {"Polygon", "MultiLineString"} and isinstance(coords, list):
        for part in coords:
            line = simple_line(part)
            if line:
                lines.append(line)
    elif gtype == "MultiPolygon" and isinstance(coords, list):
        for poly in coords:
            if not isinstance(poly, list):
                continue
            for ring in poly[:2]:
                line = simple_line(ring)
                if line:
                    lines.append(line)
    return lines


def _line_bbox(line: list[list[float]]) -> tuple[float, float, float, float] | None:
    if not line:
        return None
    lats = [p[0] for p in line]
    lngs = [p[1] for p in line]
    return min(lats), max(lats), min(lngs), max(lngs)


def _bbox_intersects(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> bool:
    return not (a[1] < b[0] or a[0] > b[1] or a[3] < b[2] or a[2] > b[3])


def _draw_reference_layers(draw: Any, project: Any, bounds: tuple[float, float, float, float]) -> list[str]:
    labels: list[str] = []
    for filename, kind, stroke, fill, width, max_features in REFERENCE_LAYERS:
        path = _geojson_path(filename)
        if not path.is_file():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        features = data.get("features") if isinstance(data, dict) else []
        if not isinstance(features, list):
            continue
        drawn = 0
        for feature in features:
            if not isinstance(feature, dict) or not isinstance(feature.get("geometry"), dict):
                continue
            for line in _coord_lines(feature["geometry"]):
                bbox = _line_bbox(line)
                if not bbox or not _bbox_intersects(bbox, bounds):
                    continue
                pts = [project(p) for p in line]
                if len(pts) < 2:
                    continue
                if kind == "polygon" and fill:
                    draw.polygon(pts, outline=stroke, fill=fill)
                    draw.line(pts + [pts[0]], fill=stroke, width=width)
                else:
                    draw.line(pts, fill=stroke, width=width)
                drawn += 1
                if drawn >= max_features:
                    break
            if drawn >= max_features:
                break
        if drawn:
            label = filename.replace(".geojson", "").replace("_", " ")
            labels.append(label)
    return labels[:5]


def map_image(data: dict[str, Any], mode: str) -> Path | None:
    pts = geometry_points(data)
    if len(pts) < 2:
        return None
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    analysis = data.get("analysis", {})
    geometry_type = str(val(data, "geometry.type", "") or "").strip().lower()
    is_line = geometry_type in {"linestring", "multilinestring"}
    sets = [s for s in analysis.get("sets_cercanas", []) if isinstance(s, dict)]
    map_pts = pts + [[safe_float(s.get("lat")), safe_float(s.get("lng"))] for s in sets if s.get("lat") not in (None, "") and s.get("lng") not in (None, "")]
    source_pts = map_pts if mode == "far" else pts
    lats = [p[0] for p in source_pts]
    lngs = [p[1] for p in source_pts]
    center_lat = (min(lats) + max(lats)) / 2.0
    center_lng = (min(lngs) + max(lngs)) / 2.0
    cos_lat = max(0.15, math.cos(math.radians(center_lat)))

    def to_xy(point: list[float]) -> tuple[float, float]:
        lat, lng = point
        return (lng - center_lng) * 111.320 * cos_lat, (lat - center_lat) * 110.574

    xy = [to_xy(p) for p in source_pts]
    min_x, max_x = min(p[0] for p in xy), max(p[0] for p in xy)
    min_y, max_y = min(p[1] for p in xy), max(p[1] for p in xy)
    pad = 0.55 if mode == "far" else 0.18
    span_x = max((max_x - min_x) * (1 + 2 * pad), 1.0)
    span_y = max((max_y - min_y) * (1 + 2 * pad), 1.0)
    center_x = (min_x + max_x) / 2.0
    center_y = (min_y + max_y) / 2.0

    width, height = 1400, 800
    plot_left, plot_top, plot_right, plot_bottom = 54, 62, width - 54, height - 78
    plot_width = plot_right - plot_left
    plot_height = plot_bottom - plot_top
    plot_aspect = plot_width / plot_height
    if span_x / span_y > plot_aspect:
        span_y = span_x / plot_aspect
    else:
        span_x = span_y * plot_aspect
    min_x, max_x = center_x - span_x / 2.0, center_x + span_x / 2.0
    min_y, max_y = center_y - span_y / 2.0, center_y + span_y / 2.0

    geo_lat_span = max(max(lats) - min(lats), 0.01)
    geo_lng_span = max(max(lngs) - min(lngs), 0.01)
    geo_bounds = (
        min(lats) - geo_lat_span * pad,
        max(lats) + geo_lat_span * pad,
        min(lngs) - geo_lng_span * pad,
        max(lngs) + geo_lng_span * pad,
    )
    img = Image.new("RGB", (width, height), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 15)
        small_font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 13)
        title_font = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 20)
    except Exception:
        font = small_font = title_font = ImageFont.load_default()
    for i in range(1, 6):
        x = int(plot_left + plot_width * i / 6)
        y = int(plot_top + plot_height * i / 6)
        draw.line([(x, plot_top), (x, plot_bottom)], fill="#DCE5EF", width=1)
        draw.line([(plot_left, y), (plot_right, y)], fill="#DCE5EF", width=1)

    def project(point: list[float]) -> tuple[int, int]:
        px, py = to_xy(point)
        x = int(plot_left + (px - min_x) / max(max_x - min_x, 1e-9) * plot_width)
        y = int(plot_top + (max_y - py) / max(max_y - min_y, 1e-9) * plot_height)
        return x, y

    ref_labels = _draw_reference_layers(draw, project, geo_bounds)
    poly = [project(p) for p in pts]
    overlay = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    overlay_draw = ImageDraw.Draw(overlay)
    if len(poly) >= 3 and not is_line:
        overlay_draw.polygon(poly, outline=(31, 78, 121, 255), fill=(96, 165, 250, 80))
        overlay_draw.line(poly + [poly[0]], fill=(31, 78, 121, 255), width=4)
    elif len(poly) >= 2:
        overlay_draw.line(poly, fill=(255, 255, 255, 230), width=11)
        overlay_draw.line(poly, fill=(31, 78, 121, 255), width=6)
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(img)
    if is_line and len(poly) >= 2:
        start_x, start_y = poly[0]
        end_x, end_y = poly[-1]
        draw.ellipse((start_x - 8, start_y - 8, start_x + 8, start_y + 8), fill="#16A34A", outline="#14532D", width=2)
        draw.ellipse((end_x - 8, end_y - 8, end_x + 8, end_y + 8), fill="#DC2626", outline="#7F1D1D", width=2)
        draw.text((start_x + 11, start_y - 11), "Inicio", fill="#14532D", font=small_font)
        draw.text((end_x + 11, end_y - 11), "Fin", fill="#7F1D1D", font=small_font)
    for idx, s in enumerate(sets[:6]):
        if s.get("lat") in (None, "") or s.get("lng") in (None, ""):
            continue
        x, y = project([safe_float(s.get("lat")), safe_float(s.get("lng"))])
        if x < plot_left or x > plot_right or y < plot_top or y > plot_bottom:
            continue
        draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill="#DC2626", outline="#7F1D1D")
        draw.text((x + 10, y - 12 + (idx % 2) * 15), str(s.get("nombre") or "SET")[:26], fill="#142034", font=small_font)
    title = "Vista de planta general - Norte arriba" if mode == "far" else "Vista de planta de detalle - Norte arriba"
    draw.rectangle((0, 0, width, 50), fill="#FFFFFF")
    draw.text((18, 13), title, fill="#1F4E79", font=title_font)

    north_x, north_y = width - 88, 78
    draw.text((north_x - 7, north_y - 27), "N", fill="#142034", font=title_font)
    draw.line([(north_x, north_y + 54), (north_x, north_y)], fill="#142034", width=4)
    draw.polygon([(north_x, north_y), (north_x - 10, north_y + 20), (north_x + 10, north_y + 20)], fill="#142034")

    target_scale = max(0.1, span_x * 0.2)
    magnitude = 10 ** math.floor(math.log10(target_scale))
    scale_km = max(v * magnitude for v in (1, 2, 5) if v * magnitude <= target_scale)
    scale_px = int(scale_km / span_x * plot_width)
    scale_x, scale_y = plot_left + 16, plot_bottom - 17
    draw.line([(scale_x, scale_y), (scale_x + scale_px, scale_y)], fill="#142034", width=5)
    draw.line([(scale_x, scale_y - 6), (scale_x, scale_y + 6)], fill="#142034", width=3)
    draw.line([(scale_x + scale_px, scale_y - 6), (scale_x + scale_px, scale_y + 6)], fill="#142034", width=3)
    draw.text((scale_x, scale_y - 28), f"{scale_km:g} km", fill="#142034", font=small_font)
    center = val(analysis, "centro", {})
    center_txt = ""
    if isinstance(center, dict) and center.get("lat") not in (None, "") and center.get("lng") not in (None, ""):
        center_txt = f"Centro aprox.: {fmt(center.get('lat'), 5)}, {fmt(center.get('lng'), 5)} | "
    refs = ", ".join(ref_labels) if ref_labels else "sin cruce de capas locales en el encuadre"
    draw.rectangle((0, height - 46, width, height), fill="#FFFFFF")
    draw.text((18, height - 31), center_txt + "Referencias locales: " + refs[:150], fill="#475569", font=small_font)
    MAP_DIR.mkdir(parents=True, exist_ok=True)
    key_src = "plan-view-v3|" + geometry_type + "|" + json.dumps(pts[:80], sort_keys=True) + mode + "|".join(ref_labels)
    key = hashlib.sha1(key_src.encode("utf-8")).hexdigest()[:16]
    out = MAP_DIR / f"zona_{key}_{mode}.png"
    img.save(out)
    return out


def add_location_maps_docx(doc: Document, data: dict[str, Any]) -> None:
    paths = [(map_image(data, "far"), "Vista de planta general (Norte arriba)"), (map_image(data, "near"), "Vista de planta de detalle (Norte arriba)")]
    if not any(path for path, _ in paths):
        return
    add_heading(doc, "Mapa de ubicacion", 2)
    for path, caption in paths:
        if not path:
            continue
        p = doc.add_paragraph(caption)
        format_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, 9)
        p.paragraph_format.keep_with_next = True
        try:
            image_p = doc.add_paragraph()
            image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_p.add_run().add_picture(str(path), width=Cm(15))
        except Exception:
            continue


def append_location_maps_pdf(story: list[Any], data: dict[str, Any], styles: dict[str, ParagraphStyle]) -> None:
    paths = [(map_image(data, "far"), "Vista de planta general (Norte arriba)"), (map_image(data, "near"), "Vista de planta de detalle (Norte arriba)")]
    if not any(path for path, _ in paths):
        return
    story.append(Paragraph("Mapa de ubicacion", styles["H"]))
    for path, caption in paths:
        if not path:
            continue
        story.append(KeepTogether([
            Paragraph(escape(caption), styles["BodyJ"]),
            RLImage(str(path), width=15.0 * cm, height=8.57 * cm),
            Spacer(1, 0.25 * cm),
        ]))


def pdf_font_name(bold: bool = False) -> str:
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold_path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    try:
        if regular.is_file() and "Arial" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Arial", str(regular)))
        if bold_path.is_file() and "Arial-Bold" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold_path)))
    except Exception:
        pass
    return "Arial-Bold" if bold and "Arial-Bold" in pdfmetrics.getRegisteredFontNames() else ("Arial" if "Arial" in pdfmetrics.getRegisteredFontNames() else ("Helvetica-Bold" if bold else "Helvetica"))


def clear_doc_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def load_doc(meta: dict[str, Any]) -> Document:
    template = Path(str(meta.get("template") or ""))
    if template.is_file():
        try:
            doc = Document(str(template))
            if template_preserve_body(meta):
                apply_template_fields(doc, meta)
            else:
                clear_doc_body(doc)
            return doc
        except Exception:
            pass
    return Document()


def _fold_template_text(value: Any) -> str:
    raw = re.sub(r"\s+", " ", str(value or "").strip())
    folded = unicodedata.normalize("NFD", raw)
    folded = "".join(ch for ch in folded if unicodedata.category(ch) != "Mn")
    return folded.upper()


def _clean_index_title(value: Any) -> str:
    title = clean_ai_text(value)
    title = re.sub(r"^\s*[-*•]+\s*", "", title)
    title = re.sub(r"^\s*\d+(?:\.\d+)*\.?\s*", "", title)
    title = re.sub(r"\s+", " ", title).strip(" .:-;")
    return title[:120].strip()


def _normalize_index_item(item: Any, default_level: int = 0) -> tuple[int, str] | None:
    level = default_level
    title: Any = ""
    if isinstance(item, dict):
        title = item.get("title") or item.get("titulo") or item.get("text") or item.get("nombre")
        raw_level = item.get("level", item.get("nivel", default_level))
        try:
            level = int(raw_level)
        except Exception:
            level = default_level
    else:
        raw = str(item or "")
        match = re.match(r"^\s*(\d+(?:\.\d+)*)", raw)
        if match:
            level = min(2, max(0, match.group(1).count(".")))
        title = raw
    clean_title = _clean_index_title(title)
    if not clean_title:
        return None
    return (min(2, max(0, level)), clean_title)


def _parse_index_items(value: Any) -> list[tuple[int, str]]:
    if value in (None, ""):
        return []
    if isinstance(value, dict):
        for key in ("items", "indice", "index", "sections", "secciones"):
            parsed = _parse_index_items(value.get(key))
            if parsed:
                return parsed
        return []
    if isinstance(value, list):
        out: list[tuple[int, str]] = []
        for idx, item in enumerate(value):
            parsed = _normalize_index_item(item, 0 if idx == 0 else 1)
            if parsed:
                out.append(parsed)
        return out
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return []
        candidate = raw
        fenced = re.search(r"```(?:json)?\s*(.*?)```", candidate, flags=re.DOTALL | re.IGNORECASE)
        if fenced:
            candidate = fenced.group(1).strip()
        if not candidate.startswith(("{", "[")):
            match = re.search(r"(\{.*\}|\[.*\])", candidate, flags=re.DOTALL)
            if match:
                candidate = match.group(1)
        if candidate.startswith(("{", "[")):
            try:
                return _parse_index_items(json.loads(candidate))
            except Exception:
                pass
        out = []
        for line in raw.splitlines():
            parsed = _normalize_index_item(line)
            if parsed:
                out.append(parsed)
        return out
    return []


def _seed_index_items(data: dict[str, Any]) -> list[tuple[int, str]]:
    candidates = [
        val(data, "meta.index_items"),
        val(data, "meta.indice"),
        val(data, "analysis.ai.index_items"),
        val(data, "analysis.ai.indice"),
        val(data, "analysis.ai.index"),
    ]
    for candidate in candidates:
        parsed = _parse_index_items(candidate)
        if parsed:
            return parsed
    return []


def _split_openai_keys(value: str) -> list[str]:
    out: list[str] = []
    for part in re.split(r"[\s,;]+", value or ""):
        key = part.strip()
        if key.startswith("sk-") and key not in out:
            out.append(key)
    return out


def _docx_openai_keys() -> list[str]:
    keys: list[str] = []
    for name in ("OPENAI_API_KEYS", "OPENAI_API_KEY", "OPENAI_API_KEY_1", "OPENAI_API_KEY_2"):
        for key in _split_openai_keys(os.environ.get(name, "")):
            if key not in keys:
                keys.append(key)
    key_file = os.environ.get("GEOVISOR_OPENAI_KEY_FILE") or str(ROOT.parent / ".openai_api_key")
    try:
        path = Path(key_file)
        if path.is_file():
            for key in _split_openai_keys(path.read_text(encoding="utf-8", errors="ignore")):
                if key not in keys:
                    keys.append(key)
    except Exception:
        pass
    return keys


def _docx_ai_index_enabled() -> bool:
    flag = str(os.environ.get("GEOVISOR_DOCX_AI_INDEX", "1")).strip().lower()
    return flag not in {"0", "false", "no", "off"}


def _report_context_for_index(data: dict[str, Any], kind: str) -> dict[str, Any]:
    analysis = data.get("analysis") if isinstance(data.get("analysis"), dict) else {}
    meta = data.get("meta") if isinstance(data.get("meta"), dict) else {}
    context: dict[str, Any] = {
        "tipo_informe": kind,
        "proyecto": meta.get("project") or data.get("project"),
        "cliente": meta.get("client") or data.get("client"),
        "nombre": meta.get("name") or analysis.get("nombre"),
    }
    if kind == "set":
        sub = data.get("subestacion") if isinstance(data.get("subestacion"), dict) else {}
        context.update({
            "subestacion": sub.get("nombre"),
            "empresa": data.get("empresa_consulta") or sub.get("empresa"),
            "filas_capacidad": len(data.get("filas", [])) if isinstance(data.get("filas"), list) else 0,
            "tiene_pit": bool(val(data, "pit.matches", [])),
            "tiene_pt": bool(val(data, "pt_coes.matches", [])),
        })
    elif kind == "line":
        context.update({
            "codigo_linea": data.get("codigo_consulta"),
            "nombre_linea": data.get("nombre_consulta"),
            "empresa": data.get("empresa_consulta"),
            "mva_operativa": val(data, "operativa.carga_mva"),
            "mva_maxima": val(data, "maxima_registrada.carga_mva"),
            "tiene_pt": bool(val(data, "pt_coes.matches", [])),
        })
    else:
        context.update({
            "area_km2": analysis.get("area_km2"),
            "longitud_km": analysis.get("longitud_km"),
            "tipo_geometria": analysis.get("tipo"),
            "tramite_recomendado": analysis.get("tramite_recomendado"),
            "sets_cercanas": len(analysis.get("sets_cercanas", [])) if isinstance(analysis.get("sets_cercanas"), list) else 0,
            "concesiones": val(analysis, "indicadores.concesiones_intersectan"),
            "lineas": val(analysis, "indicadores.lineas_intersectan"),
            "tiene_restricciones": bool(val(analysis, "restricciones.items", [])),
        })
    return context


def _extract_response_text(payload: Any) -> str:
    if not isinstance(payload, dict):
        return ""
    if isinstance(payload.get("output_text"), str):
        return payload["output_text"]
    parts: list[str] = []
    outputs = payload.get("output", [])
    for output in outputs if isinstance(outputs, list) else []:
        if not isinstance(output, dict):
            continue
        contents = output.get("content", [])
        for content in contents if isinstance(contents, list) else []:
            if isinstance(content, dict):
                value = content.get("text") or content.get("output_text")
                if isinstance(value, str):
                    parts.append(value)
    return "\n".join(parts).strip()


def _try_openai_index_items(data: dict[str, Any], kind: str) -> list[tuple[int, str]]:
    seeded = _seed_index_items(data)
    if seeded:
        return seeded
    if not _docx_ai_index_enabled():
        return []
    keys = _docx_openai_keys()
    if not keys:
        return []
    model = os.environ.get("OPENAI_MODEL", "").strip() or "gpt-5.4-mini"
    context = _report_context_for_index(data, kind)
    prompt = (
        "Propón un índice técnico en español para un informe Word institucional del geovisor CVC. "
        "Devuelve solamente JSON válido con esta forma: "
        "{\"items\":[{\"level\":0,\"title\":\"RESUMEN EJECUTIVO\"},{\"level\":1,\"title\":\"Objetivo y alcances\"}]}. "
        "Usa niveles 0, 1 y 2; máximo 18 items; no incluyas numeración dentro del título; "
        "adapta el índice estrictamente al tipo de informe y no uses nombres genéricos de otra plantilla."
    )
    body = json.dumps({
        "model": model,
        "instructions": prompt,
        "input": json.dumps(context, ensure_ascii=False, sort_keys=True),
        "max_output_tokens": 700,
    }, ensure_ascii=False).encode("utf-8")
    for key in keys[:2]:
        req = urllib.request.Request(
            "https://api.openai.com/v1/responses",
            data=body,
            method="POST",
            headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(req, timeout=8) as response:
                payload = json.loads(response.read().decode("utf-8", errors="ignore"))
            parsed = _parse_index_items(_extract_response_text(payload))
            if len(parsed) >= 5:
                return parsed
        except Exception:
            continue
    return []


def fallback_index_items(kind: str, data: dict[str, Any]) -> list[tuple[int, str]]:
    if kind == "set":
        return [
            (0, "RESUMEN EJECUTIVO"),
            (1, "Objetivo y alcance del análisis"),
            (1, "Fuentes de información utilizadas"),
            (0, "DATOS GENERALES DE LA SUBESTACIÓN"),
            (1, "Identificación de la SET evaluada"),
            (1, "Empresa, código y registros disponibles"),
            (0, "DISPONIBILIDAD DE CARGA"),
            (1, "Resumen por transformador"),
            (1, "Detalle por empresa y código de transformador"),
            (2, "Capacidad, demanda y margen disponible por barra"),
            (1, "Máximos mensuales y condición de medición"),
            (0, "PLANES DE INVERSIÓN Y TRANSMISIÓN"),
            (1, "Plan de inversiones PIT asociado"),
            (1, "Plan de Transmisión COES asociado"),
            (0, "CRITERIO DE CÁLCULO Y RECOMENDACIONES"),
        ]
    if kind == "line":
        return [
            (0, "RESUMEN EJECUTIVO"),
            (1, "Objetivo y alcance del análisis"),
            (1, "Fuente de carga operativa"),
            (0, "DATOS GENERALES DE LA LÍNEA DE TRANSMISIÓN"),
            (1, "Identificación de la línea evaluada"),
            (1, "Empresa, código y registros evaluados"),
            (0, "CARGA OPERATIVA Y MÁXIMA DEMANDA"),
            (1, "Carga operativa utilizada"),
            (1, "Carga máxima registrada y fecha"),
            (1, "Último registro disponible"),
            (1, "Serie histórica de carga"),
            (0, "PLAN DE TRANSMISIÓN Y CRITERIO TÉCNICO"),
            (1, "Anteproyectos PT COES asociados"),
            (1, "Metodología de validación"),
            (0, "CONCLUSIONES Y RECOMENDACIONES"),
        ]
    if kind == "temp_line":
        return [
            (0, "RESUMEN EJECUTIVO"),
            (1, "Objetivo y alcance de la evaluación temporal"),
            (1, "Inicio, fin y longitud aproximada"),
            (0, "DATOS GENERALES DE LA RUTA KML/KMZ"),
            (1, "Tipo de línea, tensión y configuración"),
            (1, "Mapa de ubicación referencial"),
            (0, "CONCESIONES, REDES Y CAPAS CRUZADAS"),
            (1, "Concesiones vigentes y en trámite"),
            (1, "Redes MT, SED y líneas de transmisión cercanas"),
            (1, "Predios, restricciones ambientales y culturales"),
            (0, "CRITERIO TÉCNICO DE LA LÍNEA"),
            (1, "Conductor y estructura referencial"),
            (1, "Riesgos, accesos y restricciones de ruta"),
            (0, "CONCLUSIONES Y VISITA DE CAMPO"),
        ]
    if kind == "temp_polygon":
        return [
            (0, "RESUMEN EJECUTIVO"),
            (1, "Objetivo y alcance de la evaluación temporal"),
            (1, "Área, ubicación y trámite recomendado"),
            (0, "DATOS GENERALES DEL POLÍGONO KML/KMZ"),
            (1, "Mapa de ubicación referencial y cercana"),
            (1, "Demanda y viabilidad preliminar"),
            (0, "DISPONIBILIDAD ELÉCTRICA Y CONEXIÓN"),
            (1, "Subestaciones cercanas y distancia"),
            (1, "Capacidad disponible por devanado"),
            (1, "PIT y PT asociados a SET cercanas"),
            (0, "ANÁLISIS TERRITORIAL Y RESTRICCIONES"),
            (1, "Concesiones, redes y capas cruzadas"),
            (1, "Predios identificados por API"),
            (1, "Afectaciones ambientales, culturales y relieve"),
            (0, "CONCLUSIONES Y VISITA DE CAMPO"),
        ]
    return [
        (0, "RESUMEN EJECUTIVO"),
        (1, "Objetivo y alcances"),
        (1, "Componentes evaluados"),
        (0, "DATOS GENERALES DE LA ZONA EVALUADA"),
        (1, "Ubicación geográfica y demanda"),
        (1, "Mapa de ubicación referencial"),
        (0, "EVALUACIÓN ELÉCTRICA Y DISPONIBILIDAD"),
        (1, "Subestaciones cercanas y capacidad disponible"),
        (1, "Capacidad disponible por devanado"),
        (1, "PIT y PT asociados"),
        (0, "ANÁLISIS TERRITORIAL, CONCESIONES Y RELIEVE"),
        (1, "Concesiones y redes existentes"),
        (1, "Predios, demanda y restricciones"),
        (0, "MARCO NORMATIVO Y RECOMENDACIONES"),
    ]


def fallback_table_index_items(kind: str, data: dict[str, Any]) -> list[str]:
    if kind == "set":
        return [
            "Resumen general de la subestación evaluada",
            "Disponibilidad por transformador y barra",
            "Detalle por empresa y código de transformador",
            "PIT y PT asociados a la SET",
        ]
    if kind == "line":
        return [
            "Resumen de consulta de línea",
            "Carga operativa utilizada",
            "Carga máxima registrada y último registro",
            "Plan de Transmisión COES asociado",
        ]
    if kind == "temp_line":
        return [
            "Datos generales de la ruta temporal",
            "Capas cruzadas y redes cercanas",
            "Especificación técnica preliminar",
            "Restricciones y visita de campo",
        ]
    if kind == "temp_polygon":
        return [
            "Datos generales del polígono temporal",
            "SET cercanas y disponibilidad",
            "Capacidad disponible por devanado",
            "PIT, PT, predios y restricciones",
        ]
    return [
        "Datos generales de la zona evaluada",
        "Subestaciones cercanas y disponibilidad",
        "Capacidad disponible por devanado",
        "Indicadores territoriales y eléctricos",
    ]


def report_index_items(data: dict[str, Any], kind: str) -> list[tuple[int, str]]:
    ai_items = _try_openai_index_items(data, kind)
    items = ai_items or fallback_index_items(kind, data)
    return items[:24]


def _paragraph_has_numbering(paragraph) -> bool:
    try:
        return paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
    except Exception:
        return False


def _set_paragraph_level(paragraph, level: int) -> None:
    try:
        num_pr = paragraph._p.pPr.numPr
        if num_pr is None:
            return
        ilvl = num_pr.ilvl
        if ilvl is None:
            ilvl = OxmlElement("w:ilvl")
            num_pr.insert(0, ilvl)
        ilvl.set(qn("w:val"), str(min(2, max(0, level))))
    except Exception:
        pass


def _set_paragraph_text_keep_style(paragraph, content: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = content
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(content)
        format_run(run)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _find_template_index_slots(doc: Document) -> list[Any]:
    paragraphs = list(doc.paragraphs)
    start = None
    stop = None
    for idx, paragraph in enumerate(paragraphs):
        folded = _fold_template_text(paragraph.text)
        if start is None and folded == "INDICE":
            start = idx
            continue
        if start is not None and ("INDICE DE CUADROS" in folded or "DESARROLLO DEL INFORME" in folded):
            stop = idx
            break
    if start is None or stop is None:
        return []
    slots = []
    for paragraph in paragraphs[start + 1:stop]:
        raw = str(paragraph.text or "").strip()
        if not raw:
            continue
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if _paragraph_has_numbering(paragraph) or "List" in style_name or re.match(r"^\s*\d+(?:\.\d+)*", raw):
            slots.append(paragraph)
    return slots


def _find_development_marker(doc: Document):
    for paragraph in doc.paragraphs:
        folded = _fold_template_text(paragraph.text)
        if "DESARROLLO DEL INFORME" in folded or "DE AQUI PARA ABAJO" in folded:
            return paragraph
    return None


def apply_dynamic_template_index(doc: Document, data: dict[str, Any], kind: str = "zone") -> None:
    meta = data.get("meta") if isinstance(data.get("meta"), dict) else {}
    if not template_preserve_body(meta):
        return
    slots = _find_template_index_slots(doc)
    if not slots:
        return
    items = report_index_items(data, kind)
    for paragraph, (level, title) in zip(slots, items):
        _set_paragraph_text_keep_style(paragraph, title)
        _set_paragraph_level(paragraph, level)
    for paragraph in slots[len(items):]:
        _remove_paragraph(paragraph)

    marker = _find_development_marker(doc)
    if marker is None:
        return
    for idx, title in enumerate(fallback_table_index_items(kind, data), 1):
        p = marker.insert_paragraph_before(f"Cuadro {idx}. {title}")
        format_paragraph(p, size=10.5)



def pdf_cell(value: Any, style: ParagraphStyle, limit: int = 420) -> Paragraph:
    return Paragraph(escape(clean_pdf_text(value, limit)) or "-", style)


def pdf_table(rows: list[list[Any]], styles: dict[str, ParagraphStyle], widths: list[float], header: bool = True) -> Table:
    cooked = []
    for r_idx, row in enumerate(rows):
        st = styles["CellHead"] if header and r_idx == 0 else styles["Cell"]
        cooked.append([pdf_cell(item, st, 180 if header and r_idx == 0 else 420) for item in row])
    t = Table(cooked, repeatRows=1 if header else 0, colWidths=widths, splitByRow=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C8D2E1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")) if header else ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def draw_page_header(canvas, doc) -> None:
    canvas.saveState()
    if LOGO.is_file():
        try:
            img = ImageReader(str(LOGO))
            canvas.drawImage(
                img,
                doc.pagesize[0] - doc.rightMargin - 3.8 * cm,
                doc.pagesize[1] - 1.35 * cm,
                width=3.8 * cm,
                height=1.0 * cm,
                preserveAspectRatio=True,
                mask="auto",
                anchor="ne",
            )
        except Exception:
            pass
    canvas.setStrokeColor(colors.HexColor("#D8E1EF"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, doc.pagesize[1] - 1.55 * cm, doc.pagesize[0] - doc.rightMargin, doc.pagesize[1] - 1.55 * cm)
    canvas.setFont(pdf_font_name(False), 7)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.9 * cm, f"Pagina {doc.page}")
    canvas.restoreState()


def decision(analysis: dict[str, Any]) -> str:
    demand = float(analysis.get("demanda_kw_usada") or 0)
    if demand > 500:
        return "Concesión definitiva de distribución"
    return "Convenio municipal de distribución"


def set_availability_rows(analysis: dict[str, Any], limit: int = 8) -> list[list[str]]:
    rows: list[list[str]] = []
    sets = analysis.get("sets_cercanas") if isinstance(analysis.get("sets_cercanas"), list) else []
    ordered = sorted([s for s in sets if isinstance(s, dict)], key=lambda s: (safe_float(s.get("distancia_km"), 9999), text(s.get("nombre"))))
    for s in ordered[:limit]:
        if not isinstance(s, dict):
            continue
        support = s.get("mensaje_medicion") or s.get("mensaje") or s.get("archivo_quinceminutal") or s.get("fuente") or "Capas cargadas"
        rows.append([
            text(s.get("nombre")),
            text(s.get("empresa")),
            text(s.get("ubicacion"), "Dentro" if s.get("dentro") else "Fuera"),
            f"{fmt(s.get('distancia_km'), 2)} km",
            fmt(s.get("capacidad_total_mva"), 2),
            fmt(s.get("demanda_total_mva"), 2),
            fmt(s.get("disponible_total_mva"), 2),
            f"{text(s.get('estado') or s.get('status'))}. {text(support, '')}",
        ])
    if not rows:
        rows.append(["Sin SET cercana", "-", "-", "-", "-", "-", "-", "No se identificaron SET dentro del radio evaluado con las capas cargadas."])
    return rows


def pit_investment_rows(analysis: dict[str, Any], limit: int = 10) -> list[list[str]]:
    rows: list[list[str]] = []
    pit = analysis.get("pit") if isinstance(analysis.get("pit"), dict) else {}
    items = pit.get("set_rows") if isinstance(pit.get("set_rows"), list) else []
    ordered_items = sorted([item for item in items if isinstance(item, dict)], key=lambda item: (text(item.get("set")), text(item.get("elemento") or item.get("proyecto"))))
    for item in ordered_items[:limit]:
        if not isinstance(item, dict):
            continue
        spec_parts: list[str] = []
        if item.get("potencia_mva") not in (None, ""):
            spec_parts.append(f"{fmt(item.get('potencia_mva'), 2)} MVA")
        if item.get("tension_kv") not in (None, ""):
            spec_parts.append(f"{text(item.get('tension_kv'))} kV")
        date = item.get("fecha_detalle") or item.get("fecha") or item.get("anio_reprogramado") or item.get("anio_previsto") or "-"
        rows.append([
            text(item.get("set")),
            text(item.get("pit_etiqueta"), "PIT"),
            text(item.get("elemento") or item.get("proyecto")),
            text(item.get("tipo_elemento") or item.get("prioridad")),
            text(item.get("estado") or item.get("prioridad")),
            " / ".join(spec_parts) or "-",
            text(date),
        ])
    if not rows:
        rows.append([
            "Sin registro PIT",
            "-",
            "-",
            "-",
            "-",
            "-",
            text(pit.get("message"), "No se detectaron adquisiciones o elementos PIT asociados."),
        ])
    return rows


def pt_investment_rows(analysis: dict[str, Any], limit: int = 8) -> list[list[str]]:
    rows: list[list[str]] = []
    pt = analysis.get("pt_coes") if isinstance(analysis.get("pt_coes"), dict) else {}
    items = pt.get("set_rows") if isinstance(pt.get("set_rows"), list) else []
    for item in items[:limit]:
        if not isinstance(item, dict):
            continue
        spec = " / ".join([part for part in [text(item.get("tipo"), ""), f"{text(item.get('tension_kv'), '')} kV" if item.get("tension_kv") else ""] if part])
        rows.append([
            text(item.get("set")),
            text(item.get("titulo")),
            spec or "-",
            text(item.get("departamentos")),
            fmt(item.get("score"), 1),
            text(item.get("razones") or item.get("criterio")),
        ])
    if not rows:
        rows.append(["Sin registro PT", "-", "-", "-", "-", text(pt.get("message"), "No se detectaron anteproyectos PT COES asociados.")])
    return rows


def winding_rows(analysis: dict[str, Any], limit: int = 18) -> list[list[str]]:
    rows: list[list[str]] = []
    sets = analysis.get("sets_cercanas") if isinstance(analysis.get("sets_cercanas"), list) else []
    ordered = sorted([s for s in sets if isinstance(s, dict)], key=lambda s: (safe_float(s.get("distancia_km"), 9999), text(s.get("nombre"))))
    for set_item in ordered:
        for row in (set_item.get("devanados") if isinstance(set_item.get("devanados"), list) else [])[:8]:
            if not isinstance(row, dict):
                continue
            rows.append([
                text(set_item.get("nombre")),
                text(row.get("transformador")),
                text(row.get("barra")),
                fmt(row.get("tension_kv"), 1),
                fmt(row.get("capacidad_mva"), 2),
                fmt(row.get("demanda_max_mva"), 2),
                fmt(row.get("disponible_mva"), 2),
                text(row.get("estado") or set_item.get("alerta_229")),
            ])
            if len(rows) >= limit:
                return rows
    if not rows:
        rows.append(["Sin detalle", "-", "-", "-", "-", "-", "-", "No se encontro detalle por devanado en la data disponible."])
    return rows


def predio_rows(analysis: dict[str, Any], limit: int = 8) -> list[list[str]]:
    predios = analysis.get("predios_api") if isinstance(analysis.get("predios_api"), dict) else {}
    records = predios.get("registros") if isinstance(predios.get("registros"), list) else []
    rows = [
        ["Predios identificados", fmt(predios.get("predios_identificados"), 0)],
        ["Area con predio referencial", f"{fmt(predios.get('porcentaje_area_con_predio'), 1)}% / {fmt(predios.get('area_con_predio_ha'), 1)} ha"],
        ["Propietarios/titulares muestra", ", ".join(predios.get("propietarios", [])[:6]) if isinstance(predios.get("propietarios"), list) and predios.get("propietarios") else "-"],
    ]
    for item in records[:limit]:
        if not isinstance(item, dict):
            continue
        rows.append([text(item.get("titular") or item.get("nombre")), text(item.get("detalle") or item.get("fuente"))])
    return rows


def restriction_rows(analysis: dict[str, Any]) -> list[list[str]]:
    restrictions = analysis.get("restricciones") if isinstance(analysis.get("restricciones"), dict) else {}
    items = restrictions.get("items") if isinstance(restrictions.get("items"), list) else []
    rows: list[list[str]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        rows.append([text(item.get("tipo")), text(item.get("nivel")), text(item.get("detalle"))])
    if not rows:
        rows.append(["Sin restriccion relevante", "bajo", text(restrictions.get("comentario"), "No se detectaron restricciones con las capas cargadas.")])
    return rows


def demand_profile_rows(analysis: dict[str, Any]) -> list[list[str]]:
    demand = analysis.get("demanda_perfil") if isinstance(analysis.get("demanda_perfil"), dict) else {}
    samples = demand.get("muestras") if isinstance(demand.get("muestras"), list) else []
    tipo = demand.get("tipo_probable") or demand.get("tipo_principal") or "No determinado con capas cargadas"
    rows = [
        ["Tipo probable", text(tipo)],
        ["Clientes/evidencias identificadas", fmt(demand.get("clientes_identificados"), 0)],
        ["Sustento", text(demand.get("comentario"))],
    ]
    for item in samples[:6]:
        if isinstance(item, dict):
            rows.append([text(item.get("nombre") or item.get("tipo")), text(item.get("detalle") or item.get("fuente"))])
    return rows


def ceraunico_rows(analysis: dict[str, Any]) -> list[list[str]]:
    da = analysis.get("descargas_atmosfericas") if isinstance(analysis.get("descargas_atmosfericas"), dict) else {}
    if (da.get("status") or "") != "success":
        return [["Estado", text(da.get("message"), "No se pudo consultar CERAUNICO para la geometria evaluada.")]]
    rows = [
        ["Fuente", text(da.get("fuente"), "OSINERGMIN CERAUNICO")],
        ["Año", text(da.get("anio"))],
        ["Unidad", text(da.get("unidad"), "rayos totales por km2/año")],
        ["Muestras evaluadas", fmt(da.get("muestras"), 0)],
        ["Densidad minima", fmt(da.get("min"), 2)],
        ["Densidad promedio", fmt(da.get("promedio"), 2)],
        ["Densidad maxima", fmt(da.get("max"), 2)],
        ["Clase maxima", text(da.get("clase_maxima"))],
        ["Nivel referencial", text(da.get("nivel"))],
        ["Comentario tecnico", text(da.get("comentario"))],
    ]
    classes = da.get("clases") if isinstance(da.get("clases"), list) else []
    if classes:
        summary = []
        for item in classes:
            if not isinstance(item, dict):
                continue
            summary.append(
                f"{text(item.get('clase'))}: {fmt(item.get('porcentaje'), 1)}% ({text(item.get('nivel'))})"
            )
        if summary:
            rows.append(["Distribucion por clase", "; ".join(summary)])
    return rows


def network_rows(analysis: dict[str, Any]) -> list[list[str]]:
    net = analysis.get("redes_distribucion") if isinstance(analysis.get("redes_distribucion"), dict) else {}
    return [
        ["MT intersectada", fmt(net.get("mt_intersecta"), 0)],
        ["SED dentro/cercanas", fmt(net.get("sed_dentro") or net.get("sed_cercanas"), 0)],
        ["LT intersectada", fmt(net.get("lt_intersecta"), 0)],
        ["Muestras MT", ", ".join(net.get("muestras_mt", [])[:6]) if isinstance(net.get("muestras_mt"), list) and net.get("muestras_mt") else "-"],
        ["Comentario", text(net.get("comentario"))],
    ]


def field_visit_rows(analysis: dict[str, Any]) -> list[list[str]]:
    visit = analysis.get("visita_campo") if isinstance(analysis.get("visita_campo"), dict) else {}
    activities = visit.get("actividades") if isinstance(visit.get("actividades"), list) else []
    return [
        ["Dias recomendados", fmt(visit.get("dias_recomendados"), 0)],
        ["Ratio aplicado", text(visit.get("ratio"))],
        ["Actividades", "; ".join(str(x) for x in activities[:5]) if activities else "-"],
    ]


def set_cell(cell, fill: str | None = None, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            format_run(run, size=8.7, bold=bold)


def add_heading(doc: Document, title: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = BLUE if level <= 2 else DARK


def keep_table_row_together(row: Any, repeat_header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header and tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for k, v in rows:
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        cells[0].text = k
        cells[1].text = v
        set_cell(cells[0], FILL, True)
        set_cell(cells[1])


def add_data_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        set_cell(table.rows[0].cells[i], FILL, True)
    keep_table_row_together(table.rows[0], repeat_header=True)
    for row in rows or [["Sin informacion disponible"] + [""] * (len(headers) - 1)]:
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        for i, item in enumerate(row[: len(headers)]):
            cells[i].text = item
            set_cell(cells[i])


def style_doc(doc: Document, project: str, preserve_header: bool = False) -> None:
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(3)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
        st = doc.styles[level]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLUE
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)

    sec = doc.sections[0]
    if not preserve_header:
        header = sec.header.paragraphs[0]
        header.text = f"Resumen Ejecutivo: Proyecto \"{project}\""
        if header.runs:
            format_run(header.runs[0], size=8.5, color=MUTED)
        footer = sec.footer.paragraphs[0]
        footer.text = "Pagina "
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if footer.runs:
            format_run(footer.runs[0], size=8, color=MUTED)
    else:
        for section in doc.sections:
            for container in (section.header, section.footer):
                for p in container.paragraphs:
                    for run in p.runs:
                        format_run(run, size=8.5, color=MUTED)


def build_docx(data: dict[str, Any], out: Path) -> None:
    analysis = data["analysis"]
    meta = data.get("meta", {})
    project = meta.get("project") or "Evaluación preliminar de zona eléctrica"
    author = meta.get("author") or "CVC Energía"

    doc = load_doc(meta)
    kind = str(meta.get("report_kind") or "zone")
    apply_dynamic_template_index(doc, data, kind)
    preserve_template = template_active(meta) and template_preserve_body(meta)
    style_doc(doc, project, preserve_header=template_active(meta))
    preserved_template_children = set(doc._body._element) if preserve_template else set()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(100)
    r = p.add_run("RESUMEN EJECUTIVO")
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BLUE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(project)
    r2.font.name = "Arial"
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = DARK
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run(f"Preparado por: {author}\n{datetime.now().strftime('%B %Y')}").font.size = Pt(10)
    doc.add_page_break()

    add_heading(doc, "Indice", 1)
    for item in [
        "1. Resumen ejecutivo",
        "2. Datos generales de la zona evaluada",
        "3. Evaluación eléctrica y disponibilidad",
        "4. Análisis territorial, concesiones y relieve",
        "5. Marco normativo y ruta administrativa",
        "6. Conclusiones y recomendaciones",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()

    add_heading(doc, "Índice de cuadros", 1)
    for item in [
        "Cuadro 1. Datos generales de la zona",
        "Cuadro 2. Subestaciones cercanas y disponibilidad",
        "Cuadro 3. PIT asociado a SET cercanas",
        "Cuadro 4. Indicadores territoriales y eléctricos",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()

    if preserve_template:
        body = doc._body._element
        for child in list(body):
            if child not in preserved_template_children and child.tag != qn("w:sectPr"):
                body.remove(child)

    add_heading(doc, "1. Resumen ejecutivo", 1)
    d = float(analysis.get("demanda_kw_usada") or 0)
    v = analysis.get("viabilidad", {})
    reco = decision(analysis)
    intro = (
        f"Se evaluó una zona de {fmt(analysis.get('area_km2'), 2)} km2 con demanda sustentada "
        f"de {fmt(d, 0)} kW. Bajo el umbral de 500 kW previsto para diferenciar el trámite, "
        f"la ruta administrativa recomendada es: {reco}. La viabilidad preliminar resultó "
        f"{text(v.get('nivel'))} con puntaje {fmt(v.get('score'), 0)}/100."
    )
    para = doc.add_paragraph(intro)
    format_paragraph(para)
    if val(analysis, "ai.enabled", False) and val(analysis, "ai.texto", ""):
        add_body_paragraph(doc, clean_ai_text(val(analysis, "ai.texto"))[:2500])

    add_heading(doc, "2. Datos generales de la zona evaluada", 1)
    centro = analysis.get("centro", {})
    add_kv_table(doc, [
        ("Área evaluada", f"{fmt(analysis.get('area_km2'), 2)} km2"),
        ("Centro referencial", f"{fmt(centro.get('lat'), 6)}, {fmt(centro.get('lng'), 6)}"),
        ("Demanda usada", f"{fmt(d, 0)} kW"),
        ("Origen de demanda", text(analysis.get("demanda_origen"))),
        ("Trámite recomendado", reco),
    ])
    add_location_maps_docx(doc, data)

    add_heading(doc, "3. Evaluación eléctrica y disponibilidad", 1)
    set_rows = set_availability_rows(analysis)
    add_data_table(doc, ["SET", "Empresa", "Ubicacion", "Dist.", "Cap. MVA", "Dem. MVA", "Libre MVA", "Estado / sustento"], set_rows)
    add_heading(doc, "Capacidad disponible por devanado", 2)
    add_data_table(doc, ["SET", "Transformador", "Barra", "kV", "Cap. MVA", "Dem. MVA", "Libre MVA", "Estado"], winding_rows(analysis))
    add_heading(doc, "PIT asociado a SET cercanas", 2)
    add_data_table(doc, ["SET", "PIT / ref.", "Proyecto o elemento", "Tipo", "Estado", "MVA/kV", "Fecha"], pit_investment_rows(analysis))
    add_heading(doc, "PT COES asociado a SET cercanas", 2)
    add_data_table(doc, ["SET", "Anteproyecto", "Tipo/kV", "Departamento", "Score", "Sustento"], pt_investment_rows(analysis))

    add_heading(doc, "4. Análisis territorial, concesiones y relieve", 1)
    relief = analysis.get("relieve", {})
    ind = analysis.get("indicadores", {})
    add_kv_table(doc, [
        ("Relieve", text(relief.get("descripcion") or relief.get("message"))),
        ("Elevación min/max", f"{fmt(relief.get('min_m'), 0)} / {fmt(relief.get('max_m'), 0)} m"),
        ("SET cercanas", fmt(ind.get("sets_cercanas"), 0)),
        ("Líneas intersectadas/cercanas", fmt(ind.get("lineas_intersectan"), 0)),
        ("Concesiones o trámites intersectados", fmt(ind.get("concesiones_intersectan"), 0)),
        ("Capas externas intersectadas", fmt(ind.get("externas_intersectan"), 0)),
    ])
    add_heading(doc, "Predios identificados por API", 2)
    add_kv_table(doc, [(row[0], row[1]) for row in predio_rows(analysis)])
    add_heading(doc, "Tipo de demanda identificada", 2)
    add_kv_table(doc, [(row[0], row[1]) for row in demand_profile_rows(analysis)])
    add_heading(doc, "Restricciones y afectaciones preliminares", 2)
    add_data_table(doc, ["Tipo", "Nivel", "Detalle"], restriction_rows(analysis))
    add_heading(doc, "Descargas atmosfericas CERAUNICO", 2)
    add_kv_table(doc, [(row[0], row[1]) for row in ceraunico_rows(analysis)])
    add_heading(doc, "Redes de distribucion existentes", 2)
    add_kv_table(doc, [(row[0], row[1]) for row in network_rows(analysis)])

    add_heading(doc, "5. Marco normativo y ruta administrativa", 1)
    if d > 500:
        add_body_paragraph(doc,
            "Al superar 500 kW, corresponde orientar el expediente a una concesión definitiva de distribución. "
            "La solicitud debe sustentarse con memoria descriptiva, planos UTM WGS84, cronograma, presupuesto, "
            "instrumento ambiental, servidumbres y validación de no superposición."
        )
    else:
        add_body_paragraph(doc,
            "Al no superar 500 kW, el suministro puede evaluarse mediante convenio municipal de distribución, "
            "sujeto a validación técnica municipal, planos perimétricos, servidumbres y acuerdo de concejo."
        )

    add_heading(doc, "6. Conclusiones y recomendaciones", 1)
    for reason in analysis.get("razones", [])[:8]:
        add_body_paragraph(doc, str(reason))
    add_heading(doc, "Visita de campo recomendada", 2)
    add_kv_table(doc, [(row[0], row[1]) for row in field_visit_rows(analysis)])
    add_body_paragraph(doc,
        "Este documento es un preanálisis técnico-regulatorio. No reemplaza el estudio definitivo, "
        "la revisión legal del expediente, ni la validación oficial ante la autoridad competente."
    )
    doc.save(out)


def build_pdf(data: dict[str, Any], out: Path) -> None:
    analysis = data["analysis"]
    meta = data.get("meta", {})
    project = meta.get("project") or "Evaluación preliminar de zona eléctrica"
    story = []
    styles = getSampleStyleSheet()
    base_font = pdf_font_name(False)
    bold_font = pdf_font_name(True)
    styles.add(ParagraphStyle("CenterTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"), fontName=bold_font))
    styles.add(ParagraphStyle("H", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79"), spaceBefore=12, fontName=bold_font))
    styles.add(ParagraphStyle("BodyJ", parent=styles["BodyText"], alignment=TA_JUSTIFY, fontName=base_font, fontSize=11, leading=12.65, spaceAfter=6))
    styles.add(ParagraphStyle("Cell", parent=styles["BodyText"], fontName=base_font, fontSize=6.7, leading=8.0, wordWrap="CJK", splitLongWords=True))
    styles.add(ParagraphStyle("CellHead", parent=styles["Cell"], fontName=bold_font, textColor=colors.black, fontSize=6.8, leading=8.1))
    story.append(Paragraph("RESUMEN EJECUTIVO", styles["CenterTitle"]))
    story.append(Paragraph(project, styles["Heading2"]))
    story.append(Spacer(1, 0.4 * cm))
    d = float(analysis.get("demanda_kw_usada") or 0)
    story.append(Paragraph(f"Trámite recomendado: <b>{escape(decision(analysis))}</b>", styles["BodyJ"]))
    story.append(Paragraph(f"Demanda sustentada: <b>{fmt(d, 0)} kW</b>", styles["BodyJ"]))
    story.append(Paragraph(f"Viabilidad: <b>{text(val(analysis, 'viabilidad.nivel'))} ({fmt(val(analysis, 'viabilidad.score'), 0)}/100)</b>", styles["BodyJ"]))

    sections = [
        ("Datos generales", [
            ["Área", f"{fmt(analysis.get('area_km2'), 2)} km2"],
            ["Centro", f"{fmt(val(analysis, 'centro.lat'), 6)}, {fmt(val(analysis, 'centro.lng'), 6)}"],
            ["Origen demanda", text(analysis.get("demanda_origen"))],
        ]),
        ("Indicadores", [
            ["SET cercanas", fmt(val(analysis, "indicadores.sets_cercanas"), 0)],
            ["Líneas", fmt(val(analysis, "indicadores.lineas_intersectan"), 0)],
            ["Concesiones/trámites", fmt(val(analysis, "indicadores.concesiones_intersectan"), 0)],
            ["Relieve", text(val(analysis, "relieve.descripcion"))],
        ]),
    ]
    for title, rows in sections:
        story.append(Paragraph(title, styles["H"]))
        story.append(pdf_table(rows, styles, [4.2 * cm, 10.8 * cm], header=False))
    append_location_maps_pdf(story, data, styles)

    story.append(Paragraph("Subestaciones cercanas", styles["H"]))
    set_rows = [["SET", "Empresa", "Ubicacion", "Dist.", "Cap.", "Dem.", "Libre", "Estado / sustento"]]
    set_rows.extend(set_availability_rows(analysis))
    story.append(pdf_table(set_rows, styles, [2.0*cm, 1.7*cm, 1.4*cm, 1.1*cm, 1.1*cm, 1.1*cm, 1.1*cm, 5.5*cm], header=True))
    story.append(Paragraph("Capacidad disponible por devanado", styles["H"]))
    winding = [["SET", "Trafo", "Barra", "kV", "Cap.", "Dem.", "Libre", "Estado"]]
    winding.extend(winding_rows(analysis, 16))
    story.append(pdf_table(winding, styles, [2.0*cm, 1.5*cm, 1.4*cm, 0.8*cm, 1.0*cm, 1.0*cm, 1.0*cm, 6.3*cm], header=True))
    story.append(Paragraph("PIT asociado a SET cercanas", styles["H"]))
    pit_rows = [["SET", "PIT / ref.", "Proyecto o elemento", "Tipo", "Estado", "MVA/kV", "Fecha"]]
    pit_rows.extend(pit_investment_rows(analysis))
    story.append(pdf_table(pit_rows, styles, [1.8*cm, 1.8*cm, 3.0*cm, 1.3*cm, 1.5*cm, 1.4*cm, 4.2*cm], header=True))
    story.append(Paragraph("PT COES asociado a SET cercanas", styles["H"]))
    pt_rows = [["SET", "Anteproyecto", "Tipo/kV", "Departamento", "Score", "Sustento"]]
    pt_rows.extend(pt_investment_rows(analysis))
    story.append(pdf_table(pt_rows, styles, [1.8*cm, 3.7*cm, 1.7*cm, 1.7*cm, 1.0*cm, 5.1*cm], header=True))
    story.append(Paragraph("Predios identificados por API", styles["H"]))
    story.append(pdf_table([["Indicador", "Valor"]] + predio_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
    story.append(Paragraph("Tipo de demanda identificada", styles["H"]))
    story.append(pdf_table([["Indicador", "Valor"]] + demand_profile_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
    story.append(Paragraph("Restricciones y afectaciones preliminares", styles["H"]))
    restriction_table = [["Tipo", "Nivel", "Detalle"]]
    restriction_table.extend(restriction_rows(analysis))
    story.append(pdf_table(restriction_table, styles, [3.0*cm, 2.0*cm, 10.0*cm], header=True))
    story.append(Paragraph("Redes de distribucion existentes", styles["H"]))
    story.append(pdf_table([["Indicador", "Valor"]] + network_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
    story.append(Paragraph("Conclusiones", styles["H"]))
    for reason in analysis.get("razones", [])[:8]:
        story.append(Paragraph(f"- {escape(str(reason))}", styles["BodyJ"]))
    story.append(Paragraph("Visita de campo recomendada", styles["H"]))
    story.append(pdf_table([["Indicador", "Valor"]] + field_visit_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
    if val(analysis, "ai.enabled", False) and val(analysis, "ai.texto", ""):
        story.append(Paragraph("Analisis complementario con OpenAI", styles["H"]))
        story.append(Paragraph(escape(clean_ai_text(val(analysis, "ai.texto"))[:2200]).replace("\n", "<br/>"), styles["BodyJ"]))

    doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=3*cm, leftMargin=3*cm, topMargin=2.5*cm, bottomMargin=2.5*cm)
    doc.build(story, onFirstPage=draw_page_header, onLaterPages=draw_page_header)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-docx")
    parser.add_argument("--out-pdf")
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    if args.out_docx:
        build_docx(data, Path(args.out_docx))
    if args.out_pdf:
        build_pdf(data, Path(args.out_pdf))


if __name__ == "__main__":
    main()
