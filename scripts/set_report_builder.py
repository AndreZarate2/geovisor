#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build PDF/XLSX reports for SET load availability analysis."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from zone_report_builder import (
    add_body_paragraph,
    add_data_table,
    add_heading,
    add_kv_table,
    apply_dynamic_template_index,
    load_doc,
    style_doc,
    template_active,
)
from report_charts import transformer_availability_chart, voltage_demand_history_chart

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "img" / "cvc-logo.png"


def clean_text(value: Any, limit: int | None = None) -> str:
    text = str(value or "")
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"__([^_]+)__", r"\1", text)
    text = re.sub(r"`{1,3}", "", text)
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if limit and len(text) > limit:
        return text[: max(0, limit - 1)].rstrip() + "..."
    return text


def xlsx_value(value: Any) -> Any:
    if isinstance(value, (int, float)) or value is None:
        return value
    text = clean_text(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def fmt(value: Any, digits: int = 2, dash: str = "-") -> str:
    try:
        if value is None or value == "":
            return dash
        return f"{float(value):,.{digits}f}"
    except Exception:
        return clean_text(value) or dash


def val(obj: dict[str, Any], key: str, default: Any = "") -> Any:
    cur: Any = obj
    for part in key.split("."):
        if not isinstance(cur, dict) or part not in cur:
            return default
        cur = cur[part]
    return cur


def company_display(value: Any, fallback: str = "Sin empresa") -> str:
    raw = clean_text(value)
    upper = raw.upper()
    if not raw:
        return fallback
    if upper == "ENO" or "ELECTRONOROESTE" in upper:
        return "ENOSA"
    if "RED DE ENERGIA DEL PERU" in upper or "RED DE ENERGÍA DEL PERÚ" in upper:
        return "REP"
    return raw


def company_key(value: Any) -> str:
    text = clean_text(company_display(value, value or "Sin empresa")).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\b(s|sa|saa|sac|eirl|ltda|empresa|concesionaria)\b", " ", text)
    text = re.sub(r"\b(xlsx|xls|csv|quinceminutales|excel)\b", " ", text)
    return re.sub(r"\s+", " ", text).strip() or "sin empresa"


def company_from_file(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    return company_display(Path(text).stem, "")


def analysis_company_label(row: dict[str, Any], data: dict[str, Any]) -> str:
    sub = data.get("subestacion") if isinstance(data.get("subestacion"), dict) else {}
    label = (
        row.get("empresa_capacidad")
        or row.get("empresa_quinceminutal")
        or data.get("empresa_consulta")
        or sub.get("empresa")
        or company_from_file(row.get("archivo_quinceminutal"))
    )
    fallback = company_from_file(row.get("archivo_quinceminutal")) or data.get("empresa_consulta") or "Sin empresa"
    return company_display(label, fallback)


def natural_key(value: Any) -> list[Any]:
    parts = re.split(r"(\d+)", clean_text(value).lower())
    return [int(part) if part.isdigit() else part for part in parts]


def availability_row_key(row: dict[str, Any]) -> tuple[list[Any], list[Any], list[Any]]:
    return (
        natural_key(row.get("transformador")),
        natural_key(row.get("barra")),
        natural_key(row.get("tension_kv")),
    )


def grouped_by_company(rows: list[dict[str, Any]], data: dict[str, Any] | None = None) -> list[tuple[str, list[dict[str, Any]]]]:
    data = data or {}
    groups: dict[str, dict[str, Any]] = {}
    for row in rows:
        label = analysis_company_label(row, data)
        key = company_key(label)
        groups.setdefault(key, {"label": label, "rows": []})
        if groups[key]["label"] == "Sin empresa" and label != "Sin empresa":
            groups[key]["label"] = label
        groups[key]["rows"].append(row)
    out = [(item["label"], sorted(item["rows"], key=availability_row_key)) for item in groups.values()]
    return sorted(out, key=lambda item: natural_key(item[0]))


def paragraph(value: Any, style: ParagraphStyle, limit: int | None = None) -> Paragraph:
    text = escape(clean_text(value, limit)).replace("\n", "<br/>")
    return Paragraph(text or "-", style)


def pdf_table(rows: list[list[Any]], styles: dict[str, ParagraphStyle], widths: list[float], header: bool = True) -> Table:
    cooked = []
    for r_idx, row in enumerate(rows):
        st = styles["CellHead"] if header and r_idx == 0 else styles["Cell"]
        limit = 160 if header and r_idx == 0 else 380
        cooked.append([paragraph(cell, st, limit) for cell in row])
    table = Table(cooked, colWidths=widths, repeatRows=1 if header else 0, splitByRow=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C8D2E1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")) if header else ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def draw_page_header(canvas, doc) -> None:
    canvas.saveState()
    if LOGO.is_file():
        try:
            img = ImageReader(str(LOGO))
            canvas.drawImage(
                img,
                doc.pagesize[0] - doc.rightMargin - 3.8 * cm,
                doc.pagesize[1] - 1.35 * cm,
                width=3.8 * cm,
                height=1.0 * cm,
                preserveAspectRatio=True,
                mask="auto",
                anchor="ne",
            )
        except Exception:
            pass
    canvas.setStrokeColor(colors.HexColor("#D8E1EF"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, doc.pagesize[1] - 1.55 * cm, doc.pagesize[0] - doc.rightMargin, doc.pagesize[1] - 1.55 * cm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.85 * cm, f"Pagina {doc.page}")
    canvas.restoreState()


def build_pdf(data: dict[str, Any], out: Path) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("TitleBlue", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"), fontSize=16, leading=20))
    styles.add(ParagraphStyle("Sub", parent=styles["BodyText"], alignment=TA_CENTER, textColor=colors.HexColor("#475569"), fontSize=8.7, leading=11))
    styles.add(ParagraphStyle("H", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79"), spaceBefore=10, spaceAfter=5, fontSize=12, leading=14))
    styles.add(ParagraphStyle("BodyJ", parent=styles["BodyText"], alignment=TA_JUSTIFY, leading=12.5, fontSize=8.8))
    styles.add(ParagraphStyle("Small", parent=styles["BodyText"], leading=10, fontSize=7.8, textColor=colors.HexColor("#475569")))
    styles.add(ParagraphStyle("Cell", parent=styles["BodyText"], fontSize=6.6, leading=7.8, wordWrap="CJK", splitLongWords=True))
    styles.add(ParagraphStyle("CellHead", parent=styles["Cell"], fontName="Helvetica-Bold", textColor=colors.black, fontSize=6.7, leading=8.0))

    sub = data.get("subestacion") if isinstance(data.get("subestacion"), dict) else {}
    title = f"INFORME DE DISPONIBILIDAD DE CARGA - {clean_text(sub.get('nombre') or 'SET')}"
    generated = data.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    source = "EquiposTransformadores.xlsx + quinceminutales sincronizados P/Q + PIT + PT COES"

    story: list[Any] = [
        Paragraph(title, styles["TitleBlue"]),
        Paragraph(f"Generado: {escape(str(generated))}", styles["Sub"]),
        Paragraph(f"Fuente de cálculo: {escape(source)}", styles["Sub"]),
        Spacer(1, 0.25 * cm),
    ]

    rows = [
        ["Subestación", clean_text(sub.get("nombre")), "Código SET", clean_text(sub.get("codigo"))],
        ["Empresa", clean_text(data.get("empresa_consulta") or "-"), "Quinceminutal", clean_text(data.get("archivo_quinceminutal") or "-")],
        ["Último registro", clean_text(val(data, "ultimo_registro.fecha", "-")), "Demanda último registro MVA", fmt(val(data, "ultimo_registro.mva", None), 2)],
    ]
    story.append(Paragraph("Resumen general", styles["H"]))
    story.append(pdf_table(rows, styles, [3.2 * cm, 8.6 * cm, 3.2 * cm, 8.6 * cm], header=False))
    if data.get("mensaje_medicion"):
        story.append(Spacer(1, 0.12 * cm))
        story.append(paragraph(data.get("mensaje_medicion"), styles["Small"], 1000))
    if data.get("advertencia_tecnica"):
        story.append(Spacer(1, 0.12 * cm))
        story.append(paragraph(data.get("advertencia_tecnica"), styles["Small"], 1200))

    set_summary = data.get("resumen_set") if isinstance(data.get("resumen_set"), dict) else {}
    if set_summary:
        story.append(Paragraph("Resultado para una nueva conexion", styles["H"]))
        s_rows = [
            ["Referencia en pico primario", clean_text(set_summary.get("transformador_referencia") or "-"), "Barra / tension", f"{clean_text(set_summary.get('barra_referencia') or '-')} / {clean_text(set_summary.get('tension_referencia_kv') or '-')} kV"],
            ["Disponible en pico primario MVA", fmt(set_summary.get("disponible_referencia_mva"), 2), "Planificacion en pico primario MVA", fmt(set_summary.get("disponible_planificacion_referencia_mva"), 2)],
            ["Margen conservador del periodo MVA", fmt(set_summary.get("disponible_tecnico_mva"), 2), "Planificacion conservadora MVA", fmt(set_summary.get("disponible_planificacion_mva"), 2)],
            ["Pico entrada SET MVA", fmt(set_summary.get("pico_entrada_set_mva"), 2), "Restriccion", clean_text(set_summary.get("restriccion") or "-")],
            ["Capacidad firme N-1", fmt(set_summary.get("capacidad_firme_n_1_mva"), 2, "Pendiente"), "Criterio", clean_text(set_summary.get("criterio") or "-")],
        ]
        story.append(pdf_table(s_rows, styles, [3.5 * cm, 4.8 * cm, 3.8 * cm, 11.5 * cm], header=False))

    resumen = data.get("resumen_transformadores") if isinstance(data.get("resumen_transformadores"), list) else []
    if resumen:
        story.append(Paragraph("Resumen por transformador", styles["H"]))
        t_rows = [["Transformador", "Placa MVA", "Limite primario", "Pico primario", "Fecha pico", "Barra referencia", "Disp. en pico", "Plan. en pico", "Conservador periodo", "Estado"]]
        for item in resumen:
            t_rows.append([
                item.get("transformador"),
                item.get("potencias_nominales_mva") or "-",
                fmt(item.get("capacidad_primaria_mva") or item.get("capacidad_total_mva"), 2),
                fmt(item.get("pico_primario_mva") or item.get("demanda_total_mva"), 2),
                item.get("fecha_pico_primario") or "-",
                f"{clean_text(item.get('barra_referencia') or '-')} / {clean_text(item.get('tension_referencia_kv') or '-')} kV",
                fmt(item.get("disponible_referencia_mva"), 2),
                fmt(item.get("disponible_planificacion_referencia_mva"), 2),
                fmt(item.get("disponible_tecnico_mva") if item.get("disponible_tecnico_mva") is not None else item.get("disponible_total_mva"), 2),
                item.get("estado"),
            ])
        story.append(pdf_table(t_rows, styles, [2.1 * cm, 1.8 * cm, 1.8 * cm, 1.8 * cm, 2.7 * cm, 2.6 * cm, 1.8 * cm, 1.8 * cm, 2.2 * cm, 2.3 * cm]))

    filas = data.get("filas") if isinstance(data.get("filas"), list) else []
    monthly = data.get("consumo_mensual") if isinstance(data.get("consumo_mensual"), list) else []
    monthly_by_voltage = data.get("consumo_mensual_por_tension") if isinstance(data.get("consumo_mensual_por_tension"), list) else []
    if resumen or monthly_by_voltage:
        story.append(PageBreak())
        story.append(Paragraph("Gráficas de demanda y disponibilidad", styles["H"]))
        story.append(paragraph(
            "Cada línea representa la máxima demanda mensual medida en un nivel de tensión. Los máximos de diferentes tensiones pueden ocurrir en fechas distintas y no se suman entre sí. Los gráficos por transformador comparan capacidad de placa, demanda simultánea y margen incremental. Si una barra o un transformador no tiene medición, se conserva su capacidad y se identifica expresamente como sin medición.",
            styles["Small"],
            1300,
        ))
        story.append(Spacer(1, 0.15 * cm))
        story.append(voltage_demand_history_chart(
            monthly_by_voltage,
            f"Máxima demanda mensual por nivel de tensión - {clean_text(sub.get('nombre') or 'SET')}",
            "Cada color corresponde a un nivel de tensión medido; cálculo vectorial P+jQ sin repartir la entrada primaria.",
            width=26.0 * cm,
            height=7.8 * cm,
            max_points=24,
        ))
        story.append(Spacer(1, 0.25 * cm))
        for item in resumen:
            transformer = clean_text(item.get("transformador"))
            transformer_rows = [row for row in filas if clean_text(row.get("transformador")) == transformer]
            story.append(transformer_availability_chart(
                item,
                transformer_rows,
                width=26.0 * cm,
                height=6.1 * cm,
            ))
            story.append(Spacer(1, 0.22 * cm))

    story.append(Paragraph("Condicion simultanea en el pico primario", styles["H"]))
    story.append(paragraph(
        "Esta es la condicion resumida que se muestra en la plataforma: todas las barras de cada transformador corresponden a la misma fecha y hora de su maximo primario.",
        styles["Small"],
    ))
    for company, company_rows in grouped_by_company(filas, data):
        files = sorted({clean_text(r.get("archivo_quinceminutal")) for r in company_rows if clean_text(r.get("archivo_quinceminutal"))})
        story.append(paragraph(f"{company} - {' | '.join(files)}", styles["Small"]))
        s_rows = [["Trafo", "Barra", "kV", "Lado", "Fecha pico primario", "Pico primario", "P MW", "Q Mvar", "S MVA", "FP", "Disponible", "Planificacion", "Limita"]]
        for item in company_rows:
            s_rows.append([
                item.get("transformador"),
                item.get("barra"),
                item.get("tension_kv"),
                item.get("lado_transformador") or item.get("capacidad_nota"),
                item.get("fecha_referencia_simultanea"),
                fmt(item.get("pico_primario_mva"), 2),
                fmt(item.get("p_simultaneo_mw"), 2),
                fmt(item.get("q_simultaneo_mvar"), 2),
                fmt(item.get("demanda_simultanea_mva"), 2),
                fmt(item.get("fp_simultaneo"), 3),
                fmt(item.get("disponible_referencia_mva"), 2),
                fmt(item.get("disponible_planificacion_referencia_mva"), 2),
                item.get("restriccion_referencia"),
            ])
        story.append(pdf_table(s_rows, styles, [1.7 * cm, 1.4 * cm, 1.0 * cm, 1.7 * cm, 2.6 * cm, 1.6 * cm, 1.4 * cm, 1.4 * cm, 1.4 * cm, 1.1 * cm, 1.6 * cm, 1.7 * cm, 1.7 * cm]))

    story.append(Paragraph("Detalle tecnico del periodo completo", styles["H"]))
    story.append(paragraph(
        "Los maximos independientes y el margen local son referencias historicas. El margen conservador corresponde al menor valor encontrado al revisar todos los intervalos simultaneos del periodo.",
        styles["Small"],
    ))
    for company, company_rows in grouped_by_company(filas, data):
        story.append(paragraph(company, styles["Small"]))
        b_rows = [["Trafo", "Barra", "kV", "Cap. devanado", "Max. independiente", "Fecha max.", "Margen local", "Conservador", "Planificacion", "Fecha critica", "Limita", "Estado"]]
        for item in company_rows:
            b_rows.append([
                item.get("transformador"), item.get("barra"), item.get("tension_kv"), fmt(item.get("capacidad_mva"), 2),
                fmt(item.get("demanda_max_mva"), 2),
                item.get("fecha_maxima"),
                fmt(item.get("margen_local_mva"), 2),
                fmt(item.get("disponible_tecnico_mva"), 2),
                fmt(item.get("disponible_planificacion_mva"), 2),
                item.get("fecha_critica"),
                item.get("restriccion"),
                item.get("estado"),
            ])
        story.append(pdf_table(b_rows, styles, [1.7 * cm, 1.4 * cm, 1.0 * cm, 1.7 * cm, 1.8 * cm, 2.5 * cm, 1.6 * cm, 1.6 * cm, 1.6 * cm, 2.5 * cm, 1.6 * cm, 2.1 * cm]))

    if monthly:
        story.append(Paragraph("Maximos mensuales de entrada primaria", styles["H"]))
        m_rows = [["Periodo", "Mes", "Fecha máxima", "MVA"]]
        for item in monthly:
            m_rows.append([item.get("periodo"), item.get("mes"), item.get("fecha"), fmt(item.get("mva"), 2)])
        story.append(pdf_table(m_rows, styles, [3.2 * cm, 5.0 * cm, 5.0 * cm, 3.0 * cm]))

    pit = data.get("pit") if isinstance(data.get("pit"), dict) else {}
    pit_matches = pit.get("matches") if isinstance(pit.get("matches"), list) else []
    story.append(Paragraph("Plan de inversiones PIT", styles["H"]))
    if pit.get("comentario"):
        story.append(paragraph(pit.get("comentario"), styles["BodyJ"], 1000))
    if pit_matches:
        by_company: dict[str, list[dict[str, Any]]] = {}
        for item in pit_matches[:18]:
            by_company.setdefault(company_display(item.get("empresa") or item.get("empresa_norm")), []).append(item)
        for company, items in sorted(by_company.items()):
            story.append(paragraph(company, styles["Small"]))
            p_rows = [["PIT / referencia", "Elemento", "Instalacion", "Estado", "kV", "MVA", "Fecha servicio / suscripcion"]]
            for item in items:
                dates = " / ".join([clean_text(item.get("fecha_puesta_servicio")), clean_text(item.get("fecha_suscripcion"))]).strip(" /")
                p_rows.append([
                    item.get("pit_etiqueta") or item.get("pit_titulo") or item.get("anexo"),
                    item.get("elemento"),
                    item.get("instalacion"),
                    item.get("estado"),
                    item.get("tension_kv"),
                    fmt(item.get("potencia_mva"), 2),
                    dates or "-",
                ])
            story.append(pdf_table(p_rows, styles, [3.4 * cm, 5.0 * cm, 4.6 * cm, 3.0 * cm, 1.4 * cm, 1.6 * cm, 4.6 * cm]))
    else:
        story.append(paragraph(pit.get("message") or "Sin elementos PIT asociados.", styles["Small"]))

    pt = data.get("pt_coes") if isinstance(data.get("pt_coes"), dict) else {}
    pt_matches = pt.get("matches") if isinstance(pt.get("matches"), list) else []
    story.append(Paragraph("Plan de Transmisión COES 2027-2036", styles["H"]))
    if pt_matches:
        c_rows = [["Anteproyecto", "Código", "Tipo", "kV", "Coincidencia", "Archivo"]]
        for item in pt_matches[:10]:
            c_rows.append([
                item.get("titulo"),
                item.get("codigo"),
                item.get("tipo"),
                item.get("tension_kv"),
                fmt(item.get("score"), 0) + "/100",
                item.get("archivo"),
            ])
        story.append(pdf_table(c_rows, styles, [7.0 * cm, 2.0 * cm, 3.3 * cm, 1.5 * cm, 2.6 * cm, 7.2 * cm]))
    else:
        story.append(paragraph(pt.get("message") or "Sin anteproyectos PT COES asociados.", styles["Small"]))

    story.append(Paragraph("Metodología", styles["H"]))
    story.append(paragraph(
        "Se conservan los signos de potencia activa y reactiva y se orienta cada medidor segun el flujo predominante. "
        "La vista de la plataforma toma el maximo primario de cada transformador y muestra todas sus barras en esa misma fecha y hora. "
        "Los maximos independientes por devanado se incluyen en este informe solo como referencia. Adicionalmente se evalua cada "
        "intervalo simultaneo del periodo para informar el margen conservador, limitado por el devanado objetivo y por el primario, sumando P+jQ. "
        "Las potencias nominales de los devanados no se suman. El factor de potencia de referencia es el percentil 5 "
        "historico cuando hay datos suficientes; en caso contrario se usa 0.95 inductivo. El margen de planificacion "
        "aplica 90 por ciento del limite de placa. El resultado no constituye capacidad firme N-1 sin topologia y flujo de carga.",
        styles["Small"],
    ))

    doc = SimpleDocTemplate(
        str(out),
        pagesize=landscape(A4),
        rightMargin=1.4 * cm,
        leftMargin=1.4 * cm,
        topMargin=2.05 * cm,
        bottomMargin=1.35 * cm,
    )
    doc.build(story, onFirstPage=draw_page_header, onLaterPages=draw_page_header)


def add_sheet_rows(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="E8EEF5")
    border = Border(
        left=Side(style="thin", color="C8D2E1"),
        right=Side(style="thin", color="C8D2E1"),
        top=Side(style="thin", color="C8D2E1"),
        bottom=Side(style="thin", color="C8D2E1"),
    )
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = border
    for row in rows:
        ws.append([xlsx_value(v) for v in row])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 42)


def build_xlsx(data: dict[str, Any], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    sub = data.get("subestacion") if isinstance(data.get("subestacion"), dict) else {}
    set_summary = data.get("resumen_set") if isinstance(data.get("resumen_set"), dict) else {}
    summary_rows = [
        ["Subestación", sub.get("nombre")],
        ["Código SET", sub.get("codigo")],
        ["Empresa", data.get("empresa_consulta")],
        ["Archivo quinceminutal", data.get("archivo_quinceminutal")],
        ["Estado quinceminutal", data.get("medicion_quinceminutal")],
        ["Comentario medicion", data.get("mensaje_medicion")],
        ["Fuente capacidad principal", "EquiposTransformadores.xlsx"],
        ["Regla visible", "Carga de todas las barras en la misma fecha y hora del pico primario de cada transformador"],
        ["Transformador de referencia", set_summary.get("transformador_referencia")],
        ["Barra de referencia", set_summary.get("barra_referencia")],
        ["Tension de referencia kV", set_summary.get("tension_referencia_kv")],
        ["Disponible en pico primario MVA", set_summary.get("disponible_referencia_mva")],
        ["Planificacion en pico primario MVA", set_summary.get("disponible_planificacion_referencia_mva")],
        ["Margen conservador del periodo MVA", set_summary.get("disponible_tecnico_mva")],
        ["Planificacion conservadora MVA", set_summary.get("disponible_planificacion_mva")],
        ["Capacidad firme N-1", set_summary.get("capacidad_firme_n_1_mva") or "Pendiente"],
        ["Exclusión", "Registros 2026 no considerados"],
        ["Generado", data.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]
    add_sheet_rows(ws, ["Campo", "Valor"], summary_rows)

    ws = wb.create_sheet("Barras")
    filas = data.get("filas", [])
    ordered_rows = []
    if isinstance(filas, list):
        for _, company_rows in grouped_by_company([r for r in filas if isinstance(r, dict)], data):
            ordered_rows.extend(company_rows)
    add_sheet_rows(ws, [
        "Transformador", "Barra", "Tensión barra kV", "Lado transformador", "Capacidad MVA",
        "Pico primario MVA", "Fecha pico primario", "P simultaneo MW", "Q simultaneo Mvar", "S simultaneo MVA", "FP simultaneo",
        "Disponible en pico primario MVA", "Planificacion en pico primario MVA", "Restriccion en pico primario",
        "Demanda max independiente MVA", "Demanda max MW", "Demanda max Mvar", "Factor potencia historico", "Fecha maxima independiente", "Margen local MVA",
        "Margen conservador periodo MVA", "Planificacion conservadora MVA", "Fecha critica", "Restriccion conservadora", "FP referencia nueva carga",
        "Estado", "Calidad balance", "Empresa capacidad", "Archivo quinceminutal", "Fuente capacidad", "Nota capacidad",
    ], [[
        r.get("transformador"), r.get("barra"), r.get("tension_kv"), r.get("lado_transformador"),
        r.get("capacidad_mva"), r.get("pico_primario_mva"), r.get("fecha_referencia_simultanea"), r.get("p_simultaneo_mw"),
        r.get("q_simultaneo_mvar"), r.get("demanda_simultanea_mva"), r.get("fp_simultaneo"), r.get("disponible_referencia_mva"),
        r.get("disponible_planificacion_referencia_mva"), r.get("restriccion_referencia"), r.get("demanda_max_mva"), r.get("demanda_max_mw"),
        r.get("demanda_max_mvar"), r.get("factor_potencia"), r.get("fecha_maxima"), r.get("margen_local_mva"),
        r.get("disponible_tecnico_mva"), r.get("disponible_planificacion_mva"), r.get("fecha_critica"), r.get("restriccion"), r.get("fp_referencia_nueva_carga"),
        r.get("estado"), r.get("calidad_balance"), analysis_company_label(r, data), r.get("archivo_quinceminutal"), r.get("fuente_capacidad"), r.get("capacidad_nota"),
    ] for r in ordered_rows])

    ws = wb.create_sheet("Transformadores")
    add_sheet_rows(ws, ["Transformador", "Codigo fuente", "Potencias nominales MVA", "Limite primario MVA", "Pico primario MVA", "Fecha pico primario", "Barra referencia", "kV referencia", "Disponible en pico primario MVA", "Planificacion en pico primario MVA", "Restriccion referencia", "Margen conservador periodo MVA", "Planificacion conservadora MVA", "Mejor barra conservadora", "kV", "Restriccion conservadora", "Estado", "Criterio"], [[
        r.get("transformador"), r.get("transformador_mapa"), r.get("potencias_nominales_mva"), r.get("capacidad_primaria_mva"),
        r.get("pico_primario_mva"), r.get("fecha_pico_primario"), r.get("barra_referencia"), r.get("tension_referencia_kv"),
        r.get("disponible_referencia_mva"), r.get("disponible_planificacion_referencia_mva"), r.get("restriccion_referencia"),
        r.get("disponible_tecnico_mva"), r.get("disponible_planificacion_mva"),
        r.get("barra_recomendada"), r.get("tension_recomendada_kv"), r.get("restriccion"), r.get("estado"), r.get("criterio"),
    ] for r in data.get("resumen_transformadores", []) if isinstance(r, dict)])

    ws = wb.create_sheet("Entrada primaria mensual")
    add_sheet_rows(ws, ["Periodo", "Mes", "Fecha máxima", "MVA"], [[
        r.get("periodo"), r.get("mes"), r.get("fecha"), r.get("mva"),
    ] for r in data.get("consumo_mensual", []) if isinstance(r, dict)])

    pit = data.get("pit") if isinstance(data.get("pit"), dict) else {}
    ws = wb.create_sheet("PIT")
    add_sheet_rows(ws, ["PIT / referencia", "Hoja", "Fila", "Elemento", "Instalacion", "Empresa", "Estado", "kV", "MVA", "Fecha servicio", "Fecha suscripcion", "Coincidencia"], [[
        r.get("pit_etiqueta") or r.get("pit_titulo") or r.get("anexo"), r.get("anexo"), r.get("fila"), r.get("elemento"), r.get("instalacion"), r.get("empresa"),
        r.get("estado"), r.get("tension_kv"), r.get("potencia_mva"), r.get("fecha_puesta_servicio"),
        r.get("fecha_suscripcion"), r.get("coincidencia"),
    ] for r in pit.get("matches", []) if isinstance(r, dict)])

    pt = data.get("pt_coes") if isinstance(data.get("pt_coes"), dict) else {}
    ws = wb.create_sheet("PT COES")
    add_sheet_rows(ws, ["Titulo", "Codigo", "Tipo", "kV", "Score", "Archivo", "URL"], [[
        r.get("titulo"), r.get("codigo"), r.get("tipo"), r.get("tension_kv"), r.get("score"), r.get("archivo"), r.get("url"),
    ] for r in pt.get("matches", []) if isinstance(r, dict)])

    wb.save(out)


def docx_meta(data: dict[str, Any], project: str) -> dict[str, Any]:
    template = clean_text(data.get("template")) or str(ROOT / "uploads" / "templates" / "zona_reporte.docx")
    return {
        "template": template,
        "project": project,
        "author": clean_text(data.get("author")) or "CVC Energía",
        "generated_at": data.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "preserve_template_body": True,
        "report_kind": "set",
    }


def build_docx(data: dict[str, Any], out: Path) -> None:
    sub = data.get("subestacion") if isinstance(data.get("subestacion"), dict) else {}
    set_name = clean_text(sub.get("nombre") or "SET")
    generated = clean_text(data.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    project = f"Disponibilidad de carga - {set_name}"
    meta = docx_meta(data, project)
    doc = load_doc(meta)
    doc_data = dict(data)
    doc_data["meta"] = meta
    apply_dynamic_template_index(doc, doc_data, "set")
    style_doc(doc, project, preserve_header=template_active(meta))

    add_heading(doc, "INFORME DE DISPONIBILIDAD DE CARGA", 1)
    add_body_paragraph(doc, f"Subestacion evaluada: {set_name}. Generado: {generated}.")
    add_body_paragraph(
        doc,
        "El presente informe consolida la capacidad disponible por transformador y barra, "
        "la demanda maxima registrada en los quinceminutales disponibles, asi como referencias "
        "PIT y PT COES asociadas a la subestacion consultada.",
    )

    add_heading(doc, "1. Resumen general", 2)
    add_kv_table(doc, [
        ("Subestacion", set_name),
        ("Codigo SET", clean_text(sub.get("codigo")) or "-"),
        ("Empresa", clean_text(data.get("empresa_consulta")) or "-"),
        ("Archivo quinceminutal", clean_text(data.get("archivo_quinceminutal")) or "-"),
        ("Estado quinceminutal", clean_text(data.get("medicion_quinceminutal")) or "-"),
        ("Ultimo registro", clean_text(val(data, "ultimo_registro.fecha", "-")) or "-"),
        ("Demanda ultimo registro MVA", fmt(val(data, "ultimo_registro.mva", None), 2)),
        ("Fuente principal de capacidad", "EquiposTransformadores.xlsx"),
    ])
    if data.get("mensaje_medicion"):
        add_body_paragraph(doc, data.get("mensaje_medicion"))
    if data.get("advertencia_tecnica"):
        add_body_paragraph(doc, data.get("advertencia_tecnica"))
    set_summary = data.get("resumen_set") if isinstance(data.get("resumen_set"), dict) else {}
    if set_summary:
        add_kv_table(doc, [
            ("Transformador de referencia", clean_text(set_summary.get("transformador_referencia")) or "-"),
            ("Barra de referencia", clean_text(set_summary.get("barra_referencia")) or "-"),
            ("Tension de referencia kV", clean_text(set_summary.get("tension_referencia_kv")) or "-"),
            ("Disponible en pico primario MVA", fmt(set_summary.get("disponible_referencia_mva"), 2)),
            ("Planificacion en pico primario MVA", fmt(set_summary.get("disponible_planificacion_referencia_mva"), 2)),
            ("Margen conservador del periodo MVA", fmt(set_summary.get("disponible_tecnico_mva"), 2)),
            ("Planificacion conservadora MVA", fmt(set_summary.get("disponible_planificacion_mva"), 2)),
            ("Capacidad firme N-1", fmt(set_summary.get("capacidad_firme_n_1_mva"), 2, "Pendiente")),
        ])

    resumen = data.get("resumen_transformadores") if isinstance(data.get("resumen_transformadores"), list) else []
    add_heading(doc, "2. Resumen por transformador", 2)
    add_data_table(doc, ["Transformador", "Placa MVA", "Limite primario", "Pico primario / fecha", "Barra referencia", "Disp. en pico", "Plan. en pico", "Conservador periodo", "Estado"], [[
        clean_text(item.get("transformador")),
        clean_text(item.get("potencias_nominales_mva")) or "-",
        fmt(item.get("capacidad_primaria_mva") or item.get("capacidad_total_mva"), 2),
        f"{fmt(item.get('pico_primario_mva') or item.get('demanda_total_mva'), 2)} / {clean_text(item.get('fecha_pico_primario') or '-')}",
        f"{clean_text(item.get('barra_referencia') or '-')} / {clean_text(item.get('tension_referencia_kv') or '-')} kV",
        fmt(item.get("disponible_referencia_mva"), 2),
        fmt(item.get("disponible_planificacion_referencia_mva"), 2),
        fmt(item.get("disponible_tecnico_mva") if item.get("disponible_tecnico_mva") is not None else item.get("disponible_total_mva"), 2),
        clean_text(item.get("estado")),
    ] for item in resumen if isinstance(item, dict)] or [[
        "Sin registros", "-", "-", "-", "-", "-", "-", "-", "No se encontro resumen de transformadores para la SET consultada.",
    ]])

    filas = data.get("filas") if isinstance(data.get("filas"), list) else []
    add_heading(doc, "3. Condicion simultanea en el pico primario", 2)
    add_body_paragraph(doc, "Todas las barras de cada transformador se presentan en la misma fecha y hora de su maximo primario.")
    if filas:
        for company, company_rows in grouped_by_company([r for r in filas if isinstance(r, dict)], data):
            add_heading(doc, company, 3)
            add_data_table(doc, [
                "Trafo", "Barra / kV", "Lado", "Fecha pico primario", "Pico primario", "P / Q", "S / FP", "Disponible", "Planificacion", "Limita",
            ], [[
                clean_text(item.get("transformador")),
                f"{clean_text(item.get('barra'))} / {clean_text(item.get('tension_kv'))} kV",
                clean_text(item.get("lado_transformador") or item.get("rol_devanado")),
                clean_text(item.get("fecha_referencia_simultanea")),
                fmt(item.get("pico_primario_mva"), 2),
                f"{fmt(item.get('p_simultaneo_mw'), 2)} / {fmt(item.get('q_simultaneo_mvar'), 2)}",
                f"{fmt(item.get('demanda_simultanea_mva'), 2)} / {fmt(item.get('fp_simultaneo'), 3)}",
                fmt(item.get("disponible_referencia_mva"), 2),
                fmt(item.get("disponible_planificacion_referencia_mva"), 2),
                clean_text(item.get("restriccion_referencia")),
            ] for item in company_rows])
    else:
        add_body_paragraph(doc, "No se encontro detalle de barras para la SET consultada.")

    add_heading(doc, "4. Detalle tecnico del periodo completo", 2)
    add_body_paragraph(doc, "Los maximos independientes y el margen local son referencias historicas. El margen conservador resulta de revisar todos los intervalos simultaneos del periodo.")
    if filas:
        for company, company_rows in grouped_by_company([r for r in filas if isinstance(r, dict)], data):
            add_heading(doc, company, 3)
            add_data_table(doc, [
                "Trafo", "Barra / kV", "Max. independiente", "Fecha max.", "Margen local", "Conservador", "Planificacion", "Fecha critica", "Limita", "Estado",
            ], [[
                clean_text(item.get("transformador")),
                f"{clean_text(item.get('barra'))} / {clean_text(item.get('tension_kv'))} kV",
                fmt(item.get("demanda_max_mva"), 2), clean_text(item.get("fecha_maxima")), fmt(item.get("margen_local_mva"), 2),
                fmt(item.get("disponible_tecnico_mva"), 2), fmt(item.get("disponible_planificacion_mva"), 2),
                clean_text(item.get("fecha_critica")), clean_text(item.get("restriccion")), clean_text(item.get("estado")),
            ] for item in company_rows])

    monthly = data.get("consumo_mensual") if isinstance(data.get("consumo_mensual"), list) else []
    add_heading(doc, "5. Maximos mensuales de entrada primaria", 2)
    add_data_table(doc, ["Periodo", "Mes", "Fecha maxima", "MVA"], [[
        clean_text(item.get("periodo")),
        clean_text(item.get("mes")),
        clean_text(item.get("fecha")),
        fmt(item.get("mva"), 2),
    ] for item in monthly if isinstance(item, dict)] or [[
        "Sin medicion", "-", "-", "0.00",
    ]])

    pit = data.get("pit") if isinstance(data.get("pit"), dict) else {}
    pit_matches = pit.get("matches") if isinstance(pit.get("matches"), list) else []
    add_heading(doc, "6. Plan de inversiones PIT", 2)
    if pit.get("comentario"):
        add_body_paragraph(doc, pit.get("comentario"))
    add_data_table(doc, ["PIT / referencia", "Elemento", "Instalacion", "Estado", "kV", "MVA", "Fecha"], [[
        clean_text(item.get("pit_etiqueta") or item.get("pit_titulo") or item.get("anexo")),
        clean_text(item.get("elemento")),
        clean_text(item.get("instalacion")),
        clean_text(item.get("estado")),
        clean_text(item.get("tension_kv")),
        fmt(item.get("potencia_mva"), 2),
        clean_text(" / ".join([clean_text(item.get("fecha_puesta_servicio")), clean_text(item.get("fecha_suscripcion"))]).strip(" /")) or "-",
    ] for item in pit_matches[:24] if isinstance(item, dict)] or [[
        "Sin registro PIT", "-", "-", "-", "-", "-", clean_text(pit.get("message")) or "No se detectaron elementos PIT asociados.",
    ]])

    pt = data.get("pt_coes") if isinstance(data.get("pt_coes"), dict) else {}
    pt_matches = pt.get("matches") if isinstance(pt.get("matches"), list) else []
    add_heading(doc, "7. Plan de Transmision COES", 2)
    add_data_table(doc, ["Anteproyecto", "Codigo", "Tipo", "kV", "Coincidencia", "Archivo"], [[
        clean_text(item.get("titulo")),
        clean_text(item.get("codigo")),
        clean_text(item.get("tipo")),
        clean_text(item.get("tension_kv")),
        fmt(item.get("score"), 0) + "/100",
        clean_text(item.get("archivo")),
    ] for item in pt_matches[:16] if isinstance(item, dict)] or [[
        "Sin registro PT", "-", "-", "-", "-", clean_text(pt.get("message")) or "No se detectaron anteproyectos PT COES asociados.",
    ]])

    add_heading(doc, "8. Criterio de calculo", 2)
    add_body_paragraph(
        doc,
        "Los signos de P y Q se conservan y se orientan por medidor. La vista de la plataforma toma el maximo primario "
        "de cada transformador y presenta sus otras barras exactamente en esa misma fecha y hora. Los maximos independientes "
        "se informan en este documento como referencia. Tambien se evalua todo el periodo para obtener un margen conservador. "
        "La nueva carga queda limitada por el menor margen entre su devanado y el primario, usando suma compleja P+jQ. Las potencias "
        "de placa de distintos devanados no se suman. El margen de planificacion usa 90 por ciento de placa y la "
        "capacidad firme N-1 permanece pendiente hasta disponer de topologia y estudio de flujo de carga.",
    )
    doc.save(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-pdf")
    parser.add_argument("--out-xlsx")
    parser.add_argument("--out-docx")
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    if args.out_pdf:
        build_pdf(data, Path(args.out_pdf))
    if args.out_xlsx:
        build_xlsx(data, Path(args.out_xlsx))
    if args.out_docx:
        build_docx(data, Path(args.out_docx))


if __name__ == "__main__":
    main()
