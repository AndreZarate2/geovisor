#!/usr/bin/env python3
"""Genera PDF/DOCX para geometrías temporales KML/KMZ."""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from zone_report_builder import (
    BODY_FONT,
    BODY_SIZE,
    LINE_SPACING,
    add_body_paragraph,
    add_location_maps_docx,
    append_location_maps_pdf,
    apply_template_fields,
    apply_dynamic_template_index,
    ceraunico_rows as zone_ceraunico_rows,
    demand_profile_rows as zone_demand_profile_rows,
    field_visit_rows as zone_field_visit_rows,
    format_paragraph,
    format_run,
    network_rows as zone_network_rows,
    pdf_font_name,
    predio_rows as zone_predio_rows,
    pt_investment_rows as zone_pt_investment_rows,
    restriction_rows as zone_restriction_rows,
    template_active,
    template_preserve_body,
    winding_rows as zone_winding_rows,
)


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(20, 32, 52)
MUTED = RGBColor(84, 99, 122)
FILL = "E8EEF5"
ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "img" / "cvc-logo.png"


def val(data: dict[str, Any], path: str, default: Any = "") -> Any:
    cur: Any = data
    for part in path.split("."):
        if not isinstance(cur, dict) or part not in cur:
            return default
        cur = cur[part]
    return cur


def text(value: Any, fallback: str = "s/d") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, bool):
        return "Sí" if value else "No"
    return str(value)


def fmt(value: Any, digits: int = 2, fallback: str = "s/d") -> str:
    try:
        if value is None or value == "":
            return fallback
        return f"{float(value):,.{digits}f}"
    except Exception:
        return text(value, fallback)


def clean_pdf_text(value: Any, limit: int = 520) -> str:
    out = re.sub(r"\s+", " ", text(value, "")).strip()
    if len(out) > limit:
        return out[: max(0, limit - 1)].rstrip() + "..."
    return out


def report_title(data: dict[str, Any]) -> str:
    analysis = data.get("analysis", {})
    meta = data.get("meta", {})
    name = text(meta.get("name") or analysis.get("nombre") or "KML/KMZ temporal")
    return f"Informe técnico temporal KML/KMZ - {name}"


def rows_from_hits(title: str, rows: list[dict[str, Any]], limit: int = 10) -> list[list[str]]:
    out = [[title, "Detalle"]]
    if not rows:
        out.append(["Sin cruces detectados", "Según capas cargadas"])
        return out
    for item in rows[:limit]:
        detail = text(item.get("fuente"), "")
        if item.get("distancia_km") is not None:
            detail = f"{detail} - {fmt(item.get('distancia_km'), 2)} km".strip(" -")
        out.append([text(item.get("nombre")), detail or "Según capa cargada"])
    return out


def set_availability_rows(analysis: dict[str, Any], limit: int = 8) -> list[list[Any]]:
    rows = [["SET", "Empresa", "Ubicacion", "Dist. km", "Cap. MVA", "Dem. MVA", "Libre MVA", "Estado / sustento"]]
    sets = analysis.get("sets_cercanas") if isinstance(analysis.get("sets_cercanas"), list) else []
    if not sets:
        rows.append(["Sin SET cercana", "-", "-", "-", "-", "-", "-", "No se identificaron SET dentro del radio evaluado con las capas cargadas."])
        return rows
    for item in sets[:limit]:
        if not isinstance(item, dict):
            continue
        support = item.get("mensaje_medicion") or item.get("mensaje") or item.get("archivo_quinceminutal") or item.get("fuente") or "Capas cargadas"
        rows.append([
            text(item.get("nombre")),
            text(item.get("empresa")),
            text(item.get("ubicacion"), "Dentro" if item.get("dentro") else "Fuera"),
            fmt(item.get("distancia_km"), 2),
            fmt(item.get("capacidad_total_mva"), 2),
            fmt(item.get("demanda_total_mva"), 2),
            fmt(item.get("disponible_total_mva"), 2),
            f"{text(item.get('estado') or item.get('status'))}. {text(support, '')}",
        ])
    return rows


def pit_investment_rows(analysis: dict[str, Any], limit: int = 10) -> list[list[Any]]:
    rows = [["SET", "PIT / referencia", "Proyecto o elemento", "Tipo", "Estado", "MVA / kV", "Fecha disponible"]]
    pit = analysis.get("pit") if isinstance(analysis.get("pit"), dict) else {}
    items = pit.get("set_rows") if isinstance(pit.get("set_rows"), list) else []
    if not items:
        rows.append([
            "Sin registro PIT",
            "-",
            "-",
            "-",
            "-",
            "-",
            text(pit.get("message"), "No se detectaron adquisiciones o elementos PIT asociados."),
        ])
        return rows
    for item in items[:limit]:
        if not isinstance(item, dict):
            continue
        spec_parts = []
        if item.get("potencia_mva") not in (None, ""):
            spec_parts.append(f"{fmt(item.get('potencia_mva'), 2)} MVA")
        if item.get("tension_kv") not in (None, ""):
            spec_parts.append(f"{text(item.get('tension_kv'))} kV")
        date = item.get("fecha_detalle") or item.get("fecha") or item.get("anio_reprogramado") or item.get("anio_previsto") or "s/d"
        rows.append([
            text(item.get("set")),
            text(item.get("pit_etiqueta"), "PIT"),
            text(item.get("elemento") or item.get("proyecto")),
            text(item.get("tipo_elemento") or item.get("prioridad")),
            text(item.get("estado") or item.get("prioridad")),
            " / ".join(spec_parts) or "s/d",
            text(date),
        ])
    return rows


def clear_doc_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def load_doc(meta: dict[str, Any]) -> Document:
    template = Path(str(meta.get("template") or ""))
    if template.is_file():
        try:
            doc = Document(str(template))
            if template_preserve_body(meta):
                apply_template_fields(doc, meta)
            else:
                clear_doc_body(doc)
            return doc
        except Exception:
            pass
    return Document()


def set_cell(cell, fill: str | None = None, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            format_run(run, size=8.7, bold=bold, color=DARK)


def style_doc(doc: Document, project: str, preserve_header: bool = False) -> None:
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(3)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
        try:
            st = doc.styles[level]
        except KeyError:
            continue
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLUE
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)

    sec = doc.sections[0]
    if not preserve_header:
        header = sec.header.paragraphs[0]
        header.text = project[:120]
        if header.runs:
            format_run(header.runs[0], size=8.5, color=MUTED)
        footer = sec.footer.paragraphs[0]
        footer.text = "Página "
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if footer.runs:
            format_run(footer.runs[0], size=8, color=MUTED)
    else:
        for section in doc.sections:
            for container in (section.header, section.footer):
                for p in container.paragraphs:
                    for run in p.runs:
                        format_run(run, size=8.5, color=MUTED)


def add_heading(doc: Document, title: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    try:
        p.style = f"Heading {min(level, 3)}"
    except Exception:
        pass
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = BLUE if level <= 2 else DARK
    r.font.size = Pt(15 if level == 1 else 12)


def add_cover_docx(doc: Document, data: dict[str, Any]) -> None:
    analysis = data.get("analysis", {})
    meta = data.get("meta", {})
    project = report_title(data)
    author = text(meta.get("author"), "CVC Energía")
    generated = text(meta.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    if LOGO.is_file():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p_logo.add_run().add_picture(str(LOGO), width=Inches(1.75))
        except Exception:
            pass

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("INFORME TÉCNICO TEMPORAL KML/KMZ")
    r.font.name = "Arial"
    r.font.size = Pt(23)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(text(meta.get("name") or analysis.get("nombre")))
    r2.font.name = "Arial"
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = DARK

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(f"Preparado por: {author}\nGenerado: {generated}")
    r3.font.name = "Arial"
    r3.font.size = Pt(10)
    r3.font.color.rgb = MUTED

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run("Formato institucional activo desde Gestión de datos")
    r4.font.name = "Arial"
    r4.font.size = Pt(9)
    r4.font.italic = True
    r4.font.color.rgb = MUTED
    doc.add_page_break()


def keep_table_row_together(row: Any, repeat_header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header and tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for k, v in rows:
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        cells[0].text = text(k)
        cells[1].text = text(v)
        set_cell(cells[0], FILL, True)
        set_cell(cells[1])


def add_table_docx(doc: Document, rows: list[list[Any]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, value in enumerate(rows[0]):
        table.rows[0].cells[i].text = text(value)
        set_cell(table.rows[0].cells[i], FILL, True)
    keep_table_row_together(table.rows[0], repeat_header=True)
    for row in rows[1:] or [["Sin información disponible"] + [""] * (len(rows[0]) - 1)]:
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        for i, value in enumerate(row[: len(rows[0])]):
            cells[i].text = text(value)
            set_cell(cells[i])


def add_index_docx(doc: Document, analysis: dict[str, Any]) -> None:
    add_heading(doc, "Índice", 1)
    for item in [
        "1. Resumen ejecutivo",
        "2. Datos generales de la geometría evaluada",
        "3. Concesiones y capas cruzadas",
        "4. Criterio técnico",
        "5. Nota técnica",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()
    add_heading(doc, "Índice de cuadros", 1)
    items = [
        "Cuadro 1. Resumen de la geometría temporal",
        "Cuadro 2. Concesiones y capas cruzadas",
    ]
    if analysis.get("tipo") == "linea":
        items.append("Cuadro 3. Especificación preliminar de línea")
    else:
        items.append("Cuadro 3. SET cercanas y disponibilidad")
        items.append("Cuadro 4. PIT asociado a SET cercanas")
    for item in items:
        doc.add_paragraph(item)
    doc.add_page_break()


def build_docx(data: dict[str, Any], out: Path) -> None:
    analysis = data["analysis"]
    meta = data.get("meta", {})
    project = report_title(data)
    kind = "temp_line" if analysis.get("tipo") == "linea" else "temp_polygon"
    meta["report_kind"] = meta.get("report_kind") or kind
    doc = load_doc(meta)
    doc_data = dict(data)
    doc_data["meta"] = meta
    apply_dynamic_template_index(doc, doc_data, kind)
    preserve_template = template_active(meta) and template_preserve_body(meta)
    style_doc(doc, project, preserve_header=template_active(meta))
    preserved_template_children = set(doc._body._element) if preserve_template else set()
    add_cover_docx(doc, data)
    add_index_docx(doc, analysis)

    if preserve_template:
        body = doc._body._element
        for child in list(body):
            if child not in preserved_template_children and child.tag != qn("w:sectPr"):
                body.remove(child)

    add_heading(doc, "1. Resumen ejecutivo", 1)
    if analysis.get("tipo") == "linea":
        intro = (
            f"Se evaluó una ruta temporal KML/KMZ de {fmt(analysis.get('longitud_km'), 3)} km. "
            f"La ruta parte aproximadamente desde {text(val(analysis, 'punto_inicio.referencia'))} "
            f"y llega aproximadamente a {text(val(analysis, 'punto_fin.referencia'))}."
        )
    else:
        intro = (
            f"Se evaluó un polígono temporal KML/KMZ de {fmt(analysis.get('area_km2'), 3)} km2. "
            f"La viabilidad preliminar es {text(val(analysis, 'viabilidad.nivel'))} "
            f"con puntaje {fmt(val(analysis, 'viabilidad.score'), 0)}/100."
        )
    para = doc.add_paragraph(intro)
    format_paragraph(para)

    add_heading(doc, "2. Datos generales de la geometría evaluada", 1)
    if analysis.get("tipo") == "linea":
        add_kv_table(doc, [
            ("Tipo de línea", text(val(analysis, "especificacion.tipo_linea"))),
            ("Tensión nominal", f"{fmt(val(analysis, 'especificacion.tension_kv'), 1)} kV"),
            ("Configuración", text(val(analysis, "especificacion.terna"))),
            ("Longitud aproximada", f"{fmt(analysis.get('longitud_km'), 3)} km"),
            ("Inicio aproximado", text(val(analysis, "punto_inicio.referencia"))),
            ("Fin aproximado", text(val(analysis, "punto_fin.referencia"))),
        ])
    else:
        add_kv_table(doc, [
            ("Tipo", "Polígono"),
            ("Área aproximada", f"{fmt(analysis.get('area_km2'), 3)} km2"),
            ("Trámite recomendado", text(analysis.get("tramite_recomendado"))),
            ("Viabilidad", f"{text(val(analysis, 'viabilidad.nivel'))} / {fmt(val(analysis, 'viabilidad.score'), 0)}"),
            ("Punto inicial", text(val(analysis, "punto_inicio.referencia"))),
        ])
    add_location_maps_docx(doc, data)

    add_heading(doc, "3. Concesiones y capas cruzadas", 1)
    if analysis.get("tipo") == "linea":
        hits = analysis.get("hallazgos", {})
        for title, key in [
            ("Concesiones vigentes/en trámite", "concesiones"),
            ("Líneas de transmisión cruzadas", "lt"),
            ("Media tensión cruzada", "mt"),
            ("SED cercanas", "sed"),
            ("SET cercanas a la ruta", "sets"),
            ("Predios identificados por API", "predios"),
            ("Patrimonio/CIRA", "cultural"),
            ("Ambiental/ecosistemas", "ambiental"),
            ("Productivo", "productivo"),
            ("Capas externas cruzadas", "externas"),
        ]:
            add_heading(doc, title, 2)
            add_table_docx(doc, rows_from_hits("Elemento", hits.get(key, [])))
    else:
        ind = analysis.get("indicadores", {})
        samples = analysis.get("muestras", {})
        add_table_docx(doc, [
            ["Indicador", "Valor"],
            ["Concesiones intersectadas", fmt(ind.get("concesiones_intersectan"), 0)],
            ["Líneas intersectadas", fmt(ind.get("lineas_intersectan"), 0)],
            ["MT/SED dentro o intersectada", fmt((ind.get("mt_intersecta") or 0) + (ind.get("sed_dentro") or 0), 0)],
            ["SET cercanas", fmt(ind.get("sets_cercanas"), 0)],
            ["Muestras de concesión", ", ".join(samples.get("concesiones", [])[:6]) or "s/d"],
        ])

    if analysis.get("tipo") != "linea":
        add_heading(doc, "SET cercanas y disponibilidad", 2)
        add_table_docx(doc, set_availability_rows(analysis))
        add_heading(doc, "Capacidad disponible por devanado", 2)
        add_table_docx(doc, [["SET", "Transformador", "Barra", "kV", "Cap. MVA", "Dem. MVA", "Libre MVA", "Estado"]] + zone_winding_rows(analysis))
        add_heading(doc, "PIT asociado a SET cercanas", 2)
        add_table_docx(doc, pit_investment_rows(analysis))
        add_heading(doc, "PT COES asociado a SET cercanas", 2)
        add_table_docx(doc, [["SET", "Anteproyecto", "Tipo/kV", "Departamento", "Score", "Sustento"]] + zone_pt_investment_rows(analysis))
        add_heading(doc, "Predios identificados por API", 2)
        add_table_docx(doc, [["Indicador", "Valor"]] + zone_predio_rows(analysis))
        add_heading(doc, "Tipo de demanda identificada", 2)
        add_table_docx(doc, [["Indicador", "Valor"]] + zone_demand_profile_rows(analysis))

    add_heading(doc, "Restricciones y afectaciones preliminares", 2)
    add_table_docx(doc, [["Tipo", "Nivel", "Detalle"]] + zone_restriction_rows(analysis))
    add_heading(doc, "Descargas atmosfericas CERAUNICO", 2)
    add_table_docx(doc, [["Indicador", "Valor"]] + zone_ceraunico_rows(analysis))
    add_heading(doc, "Redes de distribucion existentes", 2)
    add_table_docx(doc, [["Indicador", "Valor"]] + zone_network_rows(analysis))
    add_heading(doc, "Visita de campo recomendada", 2)
    add_table_docx(doc, [["Indicador", "Valor"]] + zone_field_visit_rows(analysis))

    add_heading(doc, "4. Criterio técnico", 1)
    if analysis.get("tipo") == "linea":
        tech = analysis.get("tecnico", {})
        add_body_paragraph(doc, "Conductor referencial: " + text(tech.get("conductor_referencial")))
        add_body_paragraph(doc, "Estructura referencial: " + text(tech.get("estructura_referencial")))
        for risk in tech.get("riesgos", []):
            add_body_paragraph(doc, str(risk))
    else:
        for reason in analysis.get("razones", [])[:10]:
            add_body_paragraph(doc, str(reason))
        relief = analysis.get("relieve", {})
        add_body_paragraph(doc, "Relieve: " + text(relief.get("descripcion") or relief.get("message")))

    add_heading(doc, "5. Nota técnica", 1)
    p = doc.add_paragraph(
        "Este informe usa una geometría temporal cargada desde KML/KMZ. No reemplaza levantamiento topográfico, "
        "estudio de servidumbre, coordinación de protecciones, ingeniería definitiva ni revisión legal del expediente."
    )
    format_paragraph(p)
    doc.save(out)


def draw_page_header(canvas, doc) -> None:
    canvas.saveState()
    if LOGO.is_file():
        try:
            img = ImageReader(str(LOGO))
            canvas.drawImage(
                img,
                doc.pagesize[0] - doc.rightMargin - 3.8 * cm,
                doc.pagesize[1] - 1.35 * cm,
                width=3.8 * cm,
                height=1.0 * cm,
                preserveAspectRatio=True,
                mask="auto",
                anchor="ne",
            )
        except Exception:
            pass
    canvas.setStrokeColor(colors.HexColor("#D8E1EF"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, doc.pagesize[1] - 1.55 * cm, doc.pagesize[0] - doc.rightMargin, doc.pagesize[1] - 1.55 * cm)
    canvas.setFont(pdf_font_name(False), 7)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.9 * cm, f"Página {doc.page}")
    canvas.restoreState()


def pdf_cell(value: Any, style: ParagraphStyle, limit: int = 520) -> Paragraph:
    return Paragraph(escape(clean_pdf_text(value, limit)) or "s/d", style)


def pdf_table(rows: list[list[Any]], styles: dict[str, ParagraphStyle], widths: list[float], header: bool = True) -> Table:
    cooked = []
    for r_idx, row in enumerate(rows):
        st = styles["CellHead"] if header and r_idx == 0 else styles["Cell"]
        cooked.append([pdf_cell(item, st, 180 if header and r_idx == 0 else 520) for item in row])
    table = Table(cooked, colWidths=widths, repeatRows=1 if header else 0, splitByRow=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C8D2E1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")) if header else ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_pdf(data: dict[str, Any], out: Path) -> None:
    analysis = data["analysis"]
    meta = data.get("meta", {})
    project = report_title(data)
    author = text(meta.get("author"), "CVC Energía")
    generated = text(meta.get("generated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    styles = getSampleStyleSheet()
    base_font = pdf_font_name(False)
    bold_font = pdf_font_name(True)
    styles.add(ParagraphStyle("CenterTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"), fontName=bold_font, fontSize=22, leading=26))
    styles.add(ParagraphStyle("SubTitle", parent=styles["Heading2"], alignment=TA_CENTER, textColor=colors.HexColor("#142034"), fontName=bold_font, fontSize=13, leading=16))
    styles.add(ParagraphStyle("H", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79"), fontName=bold_font, spaceBefore=12))
    styles.add(ParagraphStyle("BodyJ", parent=styles["BodyText"], alignment=TA_JUSTIFY, fontName=base_font, fontSize=11, leading=12.65, spaceAfter=6))
    styles.add(ParagraphStyle("Cell", parent=styles["BodyText"], fontName=base_font, fontSize=6.8, leading=8.1, wordWrap="CJK", splitLongWords=True))
    styles.add(ParagraphStyle("CellHead", parent=styles["Cell"], fontName=bold_font, textColor=colors.black, fontSize=7.0, leading=8.3))

    story: list[Any] = [
        Spacer(1, 2.2 * cm),
        Paragraph("INFORME TÉCNICO TEMPORAL KML/KMZ", styles["CenterTitle"]),
        Paragraph(escape(text(meta.get("name") or analysis.get("nombre"))), styles["SubTitle"]),
        Spacer(1, 0.5 * cm),
        Paragraph(f"Preparado por: <b>{escape(author)}</b>", styles["BodyJ"]),
        Paragraph(f"Generado: <b>{escape(generated)}</b>", styles["BodyJ"]),
        Paragraph("Formato institucional activo desde Gestión de datos.", styles["BodyJ"]),
        PageBreak(),
        Paragraph("1. Resumen ejecutivo", styles["H"]),
    ]

    if analysis.get("tipo") == "linea":
        story.append(Paragraph(
            f"Se evaluó una ruta temporal KML/KMZ de <b>{fmt(analysis.get('longitud_km'), 3)} km</b>. "
            f"Parte aproximadamente desde {escape(text(val(analysis, 'punto_inicio.referencia')))} y llega a "
            f"{escape(text(val(analysis, 'punto_fin.referencia')))}.",
            styles["BodyJ"],
        ))
        story.append(Paragraph("2. Datos generales de la geometría evaluada", styles["H"]))
        story.append(pdf_table([
            ["Campo", "Valor"],
            ["Tipo de línea", val(analysis, "especificacion.tipo_linea")],
            ["Tensión nominal", f"{fmt(val(analysis, 'especificacion.tension_kv'), 1)} kV"],
            ["Configuración", val(analysis, "especificacion.terna")],
            ["Longitud aproximada", f"{fmt(analysis.get('longitud_km'), 3)} km"],
            ["Inicio aproximado", val(analysis, "punto_inicio.referencia")],
            ["Fin aproximado", val(analysis, "punto_fin.referencia")],
        ], styles, [4.2 * cm, 10.8 * cm]))
    else:
        story.append(Paragraph(
            f"Se evaluó un polígono temporal KML/KMZ de <b>{fmt(analysis.get('area_km2'), 3)} km2</b>. "
            f"La viabilidad preliminar es <b>{escape(text(val(analysis, 'viabilidad.nivel')))}</b> "
            f"con puntaje {fmt(val(analysis, 'viabilidad.score'), 0)}/100.",
            styles["BodyJ"],
        ))
        story.append(Paragraph("2. Datos generales de la geometría evaluada", styles["H"]))
        story.append(pdf_table([
            ["Campo", "Valor"],
            ["Tipo", "Polígono"],
            ["Área aproximada", f"{fmt(analysis.get('area_km2'), 3)} km2"],
            ["Trámite recomendado", analysis.get("tramite_recomendado")],
            ["Viabilidad", f"{text(val(analysis, 'viabilidad.nivel'))} / {fmt(val(analysis, 'viabilidad.score'), 0)}"],
            ["Punto inicial", val(analysis, "punto_inicio.referencia")],
        ], styles, [4.2 * cm, 10.8 * cm]))
    append_location_maps_pdf(story, data, styles)

    story.append(Paragraph("3. Concesiones y capas cruzadas", styles["H"]))
    if analysis.get("tipo") == "linea":
        hits = analysis.get("hallazgos", {})
        for title, key in [
            ("Concesiones vigentes/en trámite", "concesiones"),
            ("Líneas de transmisión cruzadas", "lt"),
            ("Media tensión cruzada", "mt"),
            ("SED cercanas", "sed"),
            ("SET cercanas a la ruta", "sets"),
            ("Predios identificados por API", "predios"),
            ("Patrimonio/CIRA", "cultural"),
            ("Ambiental/ecosistemas", "ambiental"),
            ("Productivo", "productivo"),
            ("Capas externas cruzadas", "externas"),
        ]:
            story.append(Paragraph(title, styles["Heading3"]))
            story.append(pdf_table(rows_from_hits("Elemento", hits.get(key, [])), styles, [6.0 * cm, 9.0 * cm]))
    else:
        ind = analysis.get("indicadores", {})
        samples = analysis.get("muestras", {})
        story.append(pdf_table([
            ["Indicador", "Valor"],
            ["Concesiones intersectadas", fmt(ind.get("concesiones_intersectan"), 0)],
            ["Líneas intersectadas", fmt(ind.get("lineas_intersectan"), 0)],
            ["MT/SED dentro o intersectada", fmt((ind.get("mt_intersecta") or 0) + (ind.get("sed_dentro") or 0), 0)],
            ["SET cercanas", fmt(ind.get("sets_cercanas"), 0)],
            ["Muestras de concesión", ", ".join(samples.get("concesiones", [])[:6]) or "s/d"],
        ], styles, [4.2 * cm, 10.8 * cm]))

    if analysis.get("tipo") != "linea":
        story.append(Paragraph("SET cercanas y disponibilidad", styles["Heading3"]))
        story.append(pdf_table(set_availability_rows(analysis), styles, [2.0 * cm, 1.7 * cm, 1.4 * cm, 1.1 * cm, 1.1 * cm, 1.1 * cm, 1.1 * cm, 5.5 * cm]))
        story.append(Paragraph("Capacidad disponible por devanado", styles["Heading3"]))
        winding = [["SET", "Trafo", "Barra", "kV", "Cap.", "Dem.", "Libre", "Estado"]]
        winding.extend(zone_winding_rows(analysis, 16))
        story.append(pdf_table(winding, styles, [2.0*cm, 1.5*cm, 1.4*cm, 0.8*cm, 1.0*cm, 1.0*cm, 1.0*cm, 6.3*cm], header=True))
        story.append(Paragraph("PIT asociado a SET cercanas", styles["Heading3"]))
        story.append(pdf_table(pit_investment_rows(analysis), styles, [1.8 * cm, 1.8 * cm, 3.0 * cm, 1.3 * cm, 1.5 * cm, 1.4 * cm, 4.2 * cm]))
        story.append(Paragraph("PT COES asociado a SET cercanas", styles["Heading3"]))
        pt_rows = [["SET", "Anteproyecto", "Tipo/kV", "Departamento", "Score", "Sustento"]]
        pt_rows.extend(zone_pt_investment_rows(analysis))
        story.append(pdf_table(pt_rows, styles, [1.8*cm, 3.7*cm, 1.7*cm, 1.7*cm, 1.0*cm, 5.1*cm], header=True))
        story.append(Paragraph("Predios identificados por API", styles["Heading3"]))
        story.append(pdf_table([["Indicador", "Valor"]] + zone_predio_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
        story.append(Paragraph("Tipo de demanda identificada", styles["Heading3"]))
        story.append(pdf_table([["Indicador", "Valor"]] + zone_demand_profile_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))

    story.append(Paragraph("Restricciones y afectaciones preliminares", styles["Heading3"]))
    restriction_table = [["Tipo", "Nivel", "Detalle"]]
    restriction_table.extend(zone_restriction_rows(analysis))
    story.append(pdf_table(restriction_table, styles, [3.0*cm, 2.0*cm, 10.0*cm], header=True))
    story.append(Paragraph("Redes de distribucion existentes", styles["Heading3"]))
    story.append(pdf_table([["Indicador", "Valor"]] + zone_network_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))
    story.append(Paragraph("Visita de campo recomendada", styles["Heading3"]))
    story.append(pdf_table([["Indicador", "Valor"]] + zone_field_visit_rows(analysis), styles, [4.2*cm, 10.8*cm], header=True))

    story.append(Paragraph("4. Criterio técnico", styles["H"]))
    if analysis.get("tipo") == "linea":
        tech = analysis.get("tecnico", {})
        story.append(Paragraph("Conductor referencial: " + escape(text(tech.get("conductor_referencial"))), styles["BodyJ"]))
        story.append(Paragraph("Estructura referencial: " + escape(text(tech.get("estructura_referencial"))), styles["BodyJ"]))
        for risk in tech.get("riesgos", []):
            story.append(Paragraph("- " + escape(str(risk)), styles["BodyJ"]))
    else:
        for reason in analysis.get("razones", [])[:10]:
            story.append(Paragraph("- " + escape(str(reason)), styles["BodyJ"]))
        relief = analysis.get("relieve", {})
        story.append(Paragraph("Relieve: " + escape(text(relief.get("descripcion") or relief.get("message"))), styles["BodyJ"]))

    story.append(Paragraph("5. Nota técnica", styles["H"]))
    story.append(Paragraph(
        "Este informe usa una geometría temporal cargada desde KML/KMZ. No reemplaza levantamiento topográfico, "
        "estudio de servidumbre, coordinación de protecciones, ingeniería definitiva ni revisión legal del expediente.",
        styles["BodyJ"],
    ))
    doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=3 * cm, leftMargin=3 * cm, topMargin=2.5 * cm, bottomMargin=2.5 * cm)
    doc.build(story, onFirstPage=draw_page_header, onLaterPages=draw_page_header)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out-pdf")
    ap.add_argument("--out-docx")
    args = ap.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    if args.out_docx:
        build_docx(data, Path(args.out_docx))
    if args.out_pdf:
        build_pdf(data, Path(args.out_pdf))


if __name__ == "__main__":
    main()
