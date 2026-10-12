#!/usr/bin/env python3
"""Lookup mensual de Anexos PIT para una subestacion."""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except Exception as exc:  # pragma: no cover
    print(json.dumps({"status": "unavailable", "message": f"openpyxl no disponible: {exc}"}, ensure_ascii=False))
    sys.exit(0)


LEGAL_SUFFIXES = (
    "S A A", "S A C", "S A", "SA", "SAC", "SAA", "LTDA", "EIRL",
    "EMPRESA DE DISTRIBUCION ELECTRICA", "EMPRESA REGIONAL DE SERVICIO PUBLICO DE ELECTRICIDAD",
)

COMPANY_ALIASES = {
    "REDESUR": "RED ELECTRICA DEL SUR",
    "RED ELECTRICA DEL SUR SA": "RED ELECTRICA DEL SUR",
    "RED ELECTRICA DEL SUR S A": "RED ELECTRICA DEL SUR",
    "ELECTROSUR SA": "ELECTROSUR",
    "ELECTROSUR S A": "ELECTROSUR",
    "ELECTRO SUR ESTE": "ELECTRO SUR ESTE",
    "ELSE": "ELECTRO SUR ESTE",
    "ENOSA": "ELECTRONOROESTE",
    "ELECTRONOROESTE SA": "ELECTRONOROESTE",
    "HIDRANDINA SA": "HIDRANDINA",
    "COELVISAC": "CONSORCIO ELECTRICO DE VILLACURI",
    "CVC ENERGIA": "CONSORCIO ELECTRICO DE VILLACURI",
    "CONSORCIO ELECTRICO DE VILLACURI SA": "CONSORCIO ELECTRICO DE VILLACURI",
    "CONSORCIO ELECTRICO DE VILLACURI S A": "CONSORCIO ELECTRICO DE VILLACURI",
    "CONSORCIO ELECTRICO DE VILLACURI SAC": "CONSORCIO ELECTRICO DE VILLACURI",
    "CONSORCIO ELECTRICO DE VILLACURI S A C": "CONSORCIO ELECTRICO DE VILLACURI",
}


def plain(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("“", '"').replace("”", '"').replace("'", " ")
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip().upper()
    return text


def compact(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", plain(value))


def norm_company(value: object) -> str:
    text = plain(value)
    text = re.sub(r"\b(CIA|COMPANIA|SOCIEDAD|EMPRESA)\b", " ", text)
    for suffix in LEGAL_SUFFIXES:
        text = re.sub(rf"\b{re.escape(suffix)}\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return COMPANY_ALIASES.get(text, text)


def norm_set(value: object) -> str:
    text = plain(value)
    text = re.sub(r"\b(S\.?E\.?|SET|SUBESTACION|SUBESTACION ELECTRICA)\b", " ", text)
    text = re.sub(r"\b(MAT|AT|MT)(/(MAT|AT|MT))*\b", " ", text)
    text = re.sub(r"\b\d+(\.\d+)?\s*/\s*\d+(\.\d+)?(\s*/\s*\d+(\.\d+)?)?\s*KV\b", " ", text)
    text = re.sub(r"\b\d+(\.\d+)?\s*KV\b", " ", text)
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def header_key(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", plain(value))


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.endswith(" 00:00:00"):
        text = text[:-9]
    return re.sub(r"\s+", " ", text)


def clean_date(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 <= float(value) <= 70000:
        try:
            return from_excel(value).strftime("%Y-%m-%d")
        except Exception:
            pass
    text = clean_text(value)
    if not text or text == "-":
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    return text


def choose_col(headers: list[str], *needles: str, avoid: tuple[str, ...] = ()) -> int | None:
    best = None
    for idx, header in enumerate(headers):
        if all(n in header for n in needles) and not any(a in header for a in avoid):
            best = idx
            break
    return best


def first_col(*values: int | None) -> int | None:
    for value in values:
        if value is not None:
            return value
    return None


def find_header(rows: list[tuple[object, ...]]) -> tuple[int, list[str]] | None:
    for ridx, row in enumerate(rows[:12]):
        keys = [header_key(v) for v in row]
        joined = " ".join(keys)
        if "INSTALACION" in joined and ("TITULAR" in joined or "ELEMENTO" in joined or "PROYECTO" in joined):
            return ridx, keys
    return None


def pit_sheet_title(rows: list[tuple[object, ...]], header_idx: int) -> str:
    for row in rows[:max(1, header_idx)]:
        text = clean_text(next((v for v in row if clean_text(v)), ""))
        if text:
            return text
    return ""


def pit_period_text(value: str) -> str:
    return re.sub(r"\s+", "", value.replace("–", "-"))


def pit_label(sheet_name: str, title: str) -> str:
    text = clean_text(title)
    plain_title = plain(text)
    match = re.search(r"20\d{2}\s*[-–]\s*20\d{2}", plain_title)
    if "PLAN DE INVERSIONES" in plain_title and match:
        return "PIT " + re.sub(r"\s+", "", match.group(0).replace("–", "-"))
    if "ACTAS PENDIENTES" in plain_title:
        match = re.search(r"PIT\s*(20\d{2}\s*[-–]\s*20\d{2})", plain_title)
        return "Actas pendientes" + (f" - PIT {pit_period_text(match.group(1))}" if match else "")
    if "DADOS DE BAJA" in plain_title:
        match = re.search(r"PIT\s*(20\d{2}\s*[-–]\s*20\d{2})", plain_title)
        return "Elementos dados de baja" + (f" - PIT {pit_period_text(match.group(1))}" if match else "")
    if "RESOLUCION DIRECTORAL" in plain_title or "RESOLUCIÓN DIRECTORAL" in text.upper():
        return "Manifestacion de interes - RD 0040-2022"
    if "DECRETO DE URGENCIA" in plain_title:
        return "Relacion de proyectos - DU 048-2021"
    return text or sheet_name


def estado_norm(raw: str) -> str:
    s = plain(raw)
    if "BAJA" in s or "RETIR" in s:
        return "Retirado / baja"
    if "EN SERVICIO" in s:
        return "En servicio"
    if "EJECUCION" in s or "OBRA FISICA" in s or "PROCESO" in s:
        return "En ejecucion de obra fisica"
    if "CONCLUIDO" in s and ("SIN SERVICIO" in s or "SIN APES" in s):
        return "Concluido sin puesta en servicio"
    if "SIN AVANCE" in s or "SIN INICIO" in s:
        return "Sin avance"
    return clean_text(raw) or "Sin estado"


def tipo_elemento(elemento: str, modulo: str) -> str:
    elem = plain(elemento)
    mod = plain(modulo)
    s = plain(f"{elemento} {modulo}")
    if "CELDA" in elem:
        return "celda"
    if "LINEA" in elem and "TRANSFORMADOR" in elem:
        return "linea"
    if mod.startswith("TP-") or " TP-" in f" {mod}" or re.search(r"\b(AUTO)?TRANSFORMADOR\b", elem) or " TRAFO " in f" {elem} ":
        return "transformador"
    if "CELDA" in s:
        return "celda"
    if "LINEA" in s or "LT-" in s:
        return "linea"
    if "COMPENSADOR" in s or "BANCO" in s:
        return "compensacion"
    return "elemento"


def latest_workbook(pit_dir: Path) -> Path | None:
    files = [p for p in pit_dir.rglob("*.xls*") if not p.name.startswith("~$")]
    if not files:
        return None
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


def safe_contains(a: str, b: str) -> bool:
    return bool(a and b and min(len(a), len(b)) >= 3 and (a in b or b in a))


def set_name_tokens(value: str) -> list[str]:
    stop = {"SUR", "NORTE", "ESTE", "OESTE", "CENTRO", "NUEVA", "NUEVO", "SET", "SE"}
    return [tok for tok in norm_set(value).split() if tok and tok not in stop]


def set_name_match(target: str, candidate: str) -> bool:
    target = norm_set(target)
    candidate = norm_set(candidate)
    if not target or not candidate:
        return False
    if target == candidate:
        return True
    if len(target) >= 4 and re.search(rf"(^| ){re.escape(target)}( |$)", candidate):
        return True
    if len(candidate) >= 4 and re.search(rf"(^| ){re.escape(candidate)}( |$)", target):
        return True
    target_tokens = set(set_name_tokens(target))
    candidate_tokens = set(set_name_tokens(candidate))
    if not target_tokens or not candidate_tokens:
        return False
    shared = target_tokens & candidate_tokens
    if len(shared) >= 2:
        return True
    # Nombres cortos pero oficiales, por ejemplo ILO, deben coincidir como palabra completa.
    if shared and any(len(tok) == 3 for tok in shared):
        return True
    return False


def parse_num(value: str) -> float | None:
    try:
        return float(value.replace(",", "."))
    except Exception:
        return None


def extract_specs(*texts: str) -> dict:
    joined = plain(" ".join(t for t in texts if t))
    potencias: list[float] = []
    tensiones: list[str] = []
    for match in re.finditer(r"((?:\d+(?:[\.,]\d+)?\s*/\s*)*\d+(?:[\.,]\d+)?)\s*MVA\b", joined):
        for part in re.split(r"\s*/\s*", match.group(1)):
            num = parse_num(part)
            if num is not None:
                potencias.append(num)
    for match in re.finditer(r"((?:\d+(?:[\.,]\d+)?\s*/\s*)+\d+(?:[\.,]\d+)?|\d+(?:[\.,]\d+)?)\s*KV\b", joined):
        value = re.sub(r"\s+", "", match.group(1).replace(",", "."))
        if value and value not in tensiones:
            tensiones.append(value)
    module_match = re.search(r"\bTP-(\d{3})[A-Z]*(?:-([0-9]{3}))?", joined)
    source = "texto"
    if module_match:
        kv = str(int(module_match.group(1)))
        if kv not in tensiones:
            tensiones.append(kv)
        if module_match.group(2):
            mva = float(int(module_match.group(2)))
            if mva not in potencias:
                potencias.append(mva)
        source = "modulo" if not potencias else "texto/modulo"
    potencias = sorted(set(round(p, 3) for p in potencias))
    return {
        "potencias_mva": potencias,
        "potencia_mva": max(potencias) if potencias else None,
        "tensiones_kv": tensiones,
        "tension_kv": tensiones[0] if tensiones else "",
        "spec_source": source if (potencias or tensiones) else "",
    }


def match_score(target_set: str, target_company: str, fields: dict, companies: list[str]) -> tuple[str, list[str]]:
    set_fields = {
        "instalacion": norm_set(fields.get("instalacion", "")),
        "elemento": norm_set(fields.get("elemento", "")),
        "elementos": norm_set(fields.get("elementos", "")),
        "proyecto": norm_set(fields.get("proyecto", "")),
        "comentario": norm_set(fields.get("comentario", "")),
    }
    company_fields = {
        "titular": norm_company(companies[0] if len(companies) > 0 else ""),
        "titular_anterior": norm_company(companies[1] if len(companies) > 1 else ""),
    }
    matched = []
    company_matches = []
    company_ok = not target_company
    for label, company in company_fields.items():
        if company and (company == target_company or safe_contains(company, target_company)):
            company_ok = True
            company_matches.append(label)
    if not target_company and any(company_fields.values()):
        company_ok = True

    if target_set and set_fields["instalacion"] == target_set:
        matched.append("instalacion")
        if company_ok:
            return "fuerte", matched + company_matches
        return "medio", matched

    for label, value in set_fields.items():
        if value and set_name_match(target_set, value):
            matched.append(label)

    if not matched:
        return "", []
    if company_ok:
        return "medio", matched + company_matches
    return "bajo", matched


def scan_workbook(path: Path, nombre: str, empresa: str, limit: int, departamento: str = "") -> dict:
    target_set = norm_set(nombre)
    target_company = norm_company(empresa)
    target_department = plain(departamento)
    wb = load_workbook(path, read_only=True, data_only=True)
    matches: list[dict] = []
    sheets_scanned = 0

    for ws in wb.worksheets:
        sheets_scanned += 1
        first_rows = []
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            first_rows.append(tuple(row))
        found = find_header(first_rows)
        if not found:
            continue
        header_offset, headers = found
        sheet_title = pit_sheet_title(first_rows, header_offset)
        sheet_label = pit_label(ws.title, sheet_title)
        header_row_num = header_offset + 1
        cols = {
            "empresa": choose_col(headers, "TITULAR", avoid=("ANTERIOR",)),
            "empresa_anterior": choose_col(headers, "TITULARANTERIOR"),
            "proyecto": choose_col(headers, "PROYECTO"),
            "elemento": first_col(choose_col(headers, "NOMBRE", "ELEMENTO"), choose_col(headers, "ELEMENTOS")),
            "elementos": choose_col(headers, "ELEMENTOS"),
            "instalacion": choose_col(headers, "INSTALACION"),
            "modulo": choose_col(headers, "MODULO"),
            "area_demanda": choose_col(headers, "AREA", "DEMANDA"),
            "estado": choose_col(headers, "ESTADO"),
            "anio_previsto": first_col(choose_col(headers, "ANIO", "PREVISTO"), choose_col(headers, "ANO", "PREVISTO")),
            "anio_reprogramado": first_col(choose_col(headers, "ANIO", "REPROGRAMADO"), choose_col(headers, "ANO", "REPROGRAMADO")),
            "fecha_reprogramada": first_col(choose_col(headers, "FECHA", "REPROGRAMADA"), choose_col(headers, "FECHAEXACTA")),
            "fecha_suscripcion": choose_col(headers, "FECHA", "SUSCRIPCION"),
            "fecha_puesta": first_col(choose_col(headers, "FECHA", "PUESTA", "SERVICIO"), choose_col(headers, "FECHA", "PUESTA")),
            "fecha_baja": choose_col(headers, "FECHA", "BAJA"),
            "comentario": choose_col(headers, "COMENTARIO"),
            "acta": choose_col(headers, "ACTA", avoid=("ELEMENTO", "FECHA", "SOLICITUD")),
            "fecha_solicitud": choose_col(headers, "FECHA", "SOLICITUD"),
        }
        fill = {"empresa": "", "instalacion": "", "proyecto": "", "elemento": ""}
        empty_streak = 0
        for excel_row, row in enumerate(ws.iter_rows(min_row=header_row_num + 1, values_only=True), start=header_row_num + 1):
            values = list(row)
            if not any(clean_text(v) for v in values):
                empty_streak += 1
                if empty_streak >= 80:
                    break
                continue
            empty_streak = 0

            def val(key: str) -> object:
                idx = cols.get(key)
                return values[idx] if idx is not None and idx < len(values) else None

            empresa_raw = clean_text(val("empresa") or val("empresa_anterior") or fill["empresa"])
            instalacion = clean_text(val("instalacion") or fill["instalacion"])
            proyecto = clean_text(val("proyecto") or fill["proyecto"])
            elemento = clean_text(val("elemento") or fill["elemento"])
            if empresa_raw:
                fill["empresa"] = empresa_raw
            if instalacion:
                fill["instalacion"] = instalacion
            if proyecto:
                fill["proyecto"] = proyecto
            if elemento:
                fill["elemento"] = elemento

            comentario = clean_text(val("comentario"))
            empresa_anterior = clean_text(val("empresa_anterior"))
            elementos = clean_text(val("elementos"))
            area_demanda = clean_text(val("area_demanda"))
            score, matched_fields = match_score(
                target_set,
                target_company,
                {
                    "instalacion": instalacion,
                    "elemento": elemento,
                    "elementos": elementos,
                    "proyecto": proyecto,
                    "comentario": comentario,
                },
                [empresa_raw, empresa_anterior],
            )
            if not score:
                continue
            if score == "bajo":
                continue
            modulo = clean_text(val("modulo"))
            estado_raw = clean_text(val("estado"))
            specs = extract_specs(elemento, elementos, proyecto, instalacion, modulo, comentario)
            item = {
                "anexo": ws.title,
                "pit_titulo": sheet_title,
                "pit_etiqueta": sheet_label,
                "fila": excel_row,
                "coincidencia": score,
                "matched_fields": matched_fields,
                "empresa": empresa_raw,
                "empresa_anterior": empresa_anterior,
                "empresa_norm": norm_company(empresa_raw),
                "departamento_consulta": target_department,
                "area_demanda": area_demanda,
                "proyecto": proyecto,
                "elemento": elemento,
                "elementos": elementos,
                "instalacion": instalacion,
                "subestacion_norm": norm_set(instalacion),
                "modulo": modulo,
                "tipo_elemento": tipo_elemento(elemento, modulo),
                "potencia_mva": specs["potencia_mva"],
                "potencias_mva": specs["potencias_mva"],
                "tension_kv": specs["tension_kv"],
                "tensiones_kv": specs["tensiones_kv"],
                "spec_source": specs["spec_source"],
                "estado": estado_norm(estado_raw),
                "estado_raw": estado_raw,
                "avance": estado_norm(estado_raw),
                "acta": clean_text(val("acta")),
                "anio_previsto": clean_text(val("anio_previsto")),
                "anio_reprogramado": clean_text(val("anio_reprogramado")),
                "fecha_reprogramada": clean_date(val("fecha_reprogramada")),
                "fecha_suscripcion": clean_date(val("fecha_suscripcion")),
                "fecha_puesta_servicio": clean_date(val("fecha_puesta")),
                "fecha_baja": clean_date(val("fecha_baja")),
                "fecha_solicitud": clean_date(val("fecha_solicitud")),
                "comentario": comentario,
            }
            matches.append(item)

    rank = {"fuerte": 0, "medio": 1, "bajo": 2}
    matches.sort(key=lambda m: (rank.get(m["coincidencia"], 9), m["anexo"], m["fila"]))
    future_transformers = [
        m for m in matches
        if m["tipo_elemento"] == "transformador"
        and "EN SERVICIO" not in plain(m["estado"])
        and "BAJA" not in plain(m["estado"])
        and "RETIR" not in plain(m["estado"])
        and (m.get("potencia_mva") is not None or m.get("tension_kv") or str(m.get("modulo", "")).upper().startswith("TP-"))
    ]
    out_matches = matches[:limit]
    summary = {}
    for m in matches:
        summary[m["estado"]] = summary.get(m["estado"], 0) + 1
    return {
        "status": "success" if matches else "no_data",
        "message": "" if matches else "No se encontraron elementos PIT asociados a la subestacion consultada.",
        "file": path.name,
        "file_mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
        "sheets_scanned": sheets_scanned,
        "total_matches": len(matches),
        "criterio": "PIT filtrado por nombre explicito de SET/instalacion y empresa cuando esta disponible; se excluyen coincidencias bajas.",
        "departamento_consulta": target_department,
        "summary": summary,
        "future_transformer": bool(future_transformers),
        "future_transformers": future_transformers[:5],
        "matches": out_matches,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--nombre", required=True)
    ap.add_argument("--empresa", default="")
    ap.add_argument("--departamento", default="")
    ap.add_argument("--pit-dir", required=True)
    ap.add_argument("--limit", type=int, default=30)
    args = ap.parse_args()

    pit_dir = Path(args.pit_dir)
    if not pit_dir.exists():
        print(json.dumps({"status": "unavailable", "message": "Carpeta PIT no configurada."}, ensure_ascii=False))
        return 0
    workbook = latest_workbook(pit_dir)
    if not workbook:
        print(json.dumps({"status": "unavailable", "message": "No hay Excel PIT cargado."}, ensure_ascii=False))
        return 0
    try:
        result = scan_workbook(workbook, args.nombre, args.empresa, args.limit, args.departamento)
    except Exception as exc:
        result = {"status": "unavailable", "message": f"No se pudo leer PIT: {exc}", "file": workbook.name}
    print(json.dumps(result, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
